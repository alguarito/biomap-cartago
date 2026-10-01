"""Población por celda de Cartago (2026): censo 2018 del DANE por manzana y por sección rural, llevado a las
proyecciones municipales del DANE; HRSL (Meta/CIESIN) y ESA WorldCover 2021 solo como apoyo dentro de manzanas grandes
y dentro de cada cuadro rural de 1 km.

FUENTES (todas abiertas; URL exactas en las constantes de abajo)
  1. DANE — Marco Geoestadístico Nacional 2018 integrado con el CNPV 2018 (variables anonimizadas):
       manzana censal   MGN2018/Serv_CapaManzanaInt_2018/MapServer/0     (PERSONAS_S = personas en hogares particulares;
                                                                         TP27_PERSO = número de personas; PERSONAS_L = en LEA)
       sección rural    MGN2018/Serv_CapaSeccionRuralInt_2018/MapServer/0 (STPERSON_S; STP27_PERS; STPERSON_L)
     Definiciones de los campos: tablas de manzana y sección rural del instructivo (copia en fuentes/poblacion/). Como los
     LEA solo se publican por clase, PERSONAS_L = STPERSON_L = 0 y TP27_PERSO = PERSONAS_S en Cartago (el script lo
     comprueba); se usan los campos de hogares particulares.
     Instructivo de uso (DANE, «MGN2018_Integrado_CNPV2018_InstructivoUso.pdf»): las manzanas o secciones rurales con
     1 a 3 viviendas ocupadas con personas presentes se UNEN a otra manzana de la misma sección urbana (o a otra sección
     del mismo sector rural); los Lugares Especiales de Alojamiento (LEA) solo se publican a nivel de clase. Los totales
     se conservan: en Cartago, Σ manzanas de la cabecera = 114.571 = «personas en hogares particulares, cabecera» de la
     ficha CNPV 2018, y Σ secciones rurales = 2.824 = «centros poblados y rural disperso» (el script lo comprueba).
  2. DANE — PPED, «Serie municipal de población por área 2018-2042» (hoja PobMunicipalxÁrea, «Actualizado el 30 de Julio
     de 2025»; base CNPV 2018). Totales de cabecera y de «centros poblados y rural disperso» del año ANIO_ANCLA.
  3. DANE — MGN 2018 (MGN2018/Serv_CapasMGN_2018: capa 305 «Zona Urbana», capa 317 «Municipio»). La cabecera del MGN
     2018 (19,1 km²) incluye a Zaragoza; el polígono de cabecera de OpenStreetMap (15,2 km²) no.
  4. DANE — Grilla estadística nivel 6 (1 km), CNPV 2018 (Grilla_DANE/Serv_Grilla_DANE, capa 1). Solo se usa para
     repartir la población rural dispersa DENTRO de cada sección rural. En la grilla no hay ninguna celda publicada con
     menos de 4 viviendas ocupadas, y 302 celdas de la ventana tienen viviendas > 0 con personas = hogares = ocupadas = 0:
     se interpretan como dato suprimido (supuesto inferido de los datos, coherente con el umbral de 4 viviendas ocupadas
     del instructivo; el DANE no documenta la regla de la grilla). Se imputan con la esperanza de las viviendas ocupadas
     condicionada a ≤ 3 (binomial con la tasa de ocupación rural de Cartago) × personas por vivienda ocupada rural.
  5. HRSL v1.5 (Meta Data for Good y CIESIN; población 2020, edificaciones detectadas en imágenes de 2016), VRT público
       https://dataforgood-fb-data.s3.amazonaws.com/hrsl-cogs/hrsl_general/hrsl_general-latest.vrt  (lectura por ventana)
     Licencia CC BY 4.0. Usos: reparto dentro de manzanas ≥ 2 ha y dentro de cada cuadro de 1 km en lo rural; versiones
     auxiliares (HRSL sin escalar y HRSL escalado por clase, que era el método prescrito) y validación.
  6. WorldPop 2020 (wpgppop, CC BY 4.0), API de estadísticas por polígono: solo contraste de totales.
  7. Ficha municipal CNPV 2018 de Cartago (DANE): control de totales por clase, hogares particulares y LEA.
  8. OpenStreetMap (ODbL): comunas, vías principales y puntos de control (comun.comunas(), comun.osm()); la zona de
     Zaragoza de comun.comunas() sale de ESA WorldCover 2021 v200 (CC BY 4.0) vía datos/zaragoza.json.
  9. ESA WorldCover 10 m 2021 v200 (CC BY 4.0; Zanaga et al. 2022, doi:10.5281/zenodo.7254221), clase 50 «construido»
     (edificios, vías y otras estructuras; PUM v2.0, tabla 3). Se lee la ventana local que guarda scripts/worldcover.py
     (rejilla común + 0,01°); este script no la descarga. Si falta, el reparto vuelve a usar solo HRSL.
  Términos de uso del DANE (1-4, 7): uso y transformación autorizados citando «Fuente: Departamento Administrativo
  Nacional de Estadística: www.dane.gov.co».

PASOS
  1. Descargas en paralelo con reintentos y caché en fuentes/poblacion/. Cada caché ArcGIS guarda la consulta (URL,
     filtro, caja) y la fecha; si la consulta cambia, se vuelve a pedir.
  2. HRSL → rejilla común conservando la masa (reparto por área de solape; contraste con rasterio Resampling.sum; ±2 %).
  3. Base censal 2018 (personas en hogares particulares):
       a) manzanas (cabecera y centros poblados): cada manzana reparte sus personas de manera uniforme sobre su polígono,
          calculado en una sub-rejilla 5 × 5 más fina; en manzanas ≥ 2 ha, solo sobre las sub-celdas con edificación
          HRSL si cubren ≥ 50 % de lo construido según WorldCover 2021 en la manzana, y si no, sobre la unión de ambas
          (supuesto documentado). Se suma a la rejilla común sin pérdida.
       b) población rural dispersa de cada sección rural = personas de la sección − personas de sus manzanas de centro
          poblado (emparejadas por código: clase 2 → 3 en la 6.ª posición); se reparte dentro de la sección (fuera de
          las zonas urbanas) en proporción a: personas de la grilla de 1 km (con las celdas suprimidas imputadas) −
          (1 + τ) × personas de TODAS las manzanas censales del cuadro (Cartago y municipios vecinos, p. ej. Puerto
          Caldas), con τ = mayor diferencia relativa grilla/manzanas en los cuadros ≥ 95 % urbanos (ruido de
          agregación); dentro del cuadro, según la unión de HRSL 2016 y WorldCover 2021 (peso del píxel = máximo de sus
          dos participaciones normalizadas en el cuadro; solo HRSL donde la ventana de WorldCover no llega). Los totales
          por sección se conservan exactamente. Celdas suprimidas imputadas con las razones de las secciones censales
          (y, como contraste, con las de las celdas rurales publicadas de la grilla).
  4. Anclaje 2026 por clase: cabecera × (DANE cabecera 2026 / 114.571); centros poblados y rural disperso ×
     (DANE resto 2026 / 2.824). Fuera del municipio (MGN) la capa queda sin dato.
  5. Publica 'poblacion' (comun.guardar_capa, escala 0,01, hab/celda) y añade estadísticas con la cabecera MGN;
     datos/series/poblacion.json; datos/vectores/poblacion.geojson (cabecera, municipio y secciones rurales).
  6. Verificaciones y validación: totales por clase; grilla vs manzanas a 1 km; error de los métodos basados en HRSL
     frente al censo por manzana a 250 m, 500 m, 1 km y por comuna; sensibilidad al supuesto dentro de manzana y a la
     regla de manzanas grandes; fuera de la cabecera, desacuerdo entre repartos con HRSL, WorldCover y su unión (por celda,
     250 m, 500 m, 1 km, cuadro de la grilla y sección rural) y prueba análoga de esas señales en la cabecera;
     conservación; PNG vs .npy; Parque Bolívar; aeródromo; río La Vieja; anillo periurbano; orientación norte/sur.

Uso: .venv/bin/python scripts/poblacion.py [--refrescar]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import comun  # noqa: E402

NOMBRE = "poblacion"
DIR = os.path.join(comun.FUENTES, "poblacion")
os.makedirs(DIR, exist_ok=True)
UA = {"User-Agent": "Mozilla/5.0 (BioMap Cartago; datos abiertos)"}

URL_VRT = "https://dataforgood-fb-data.s3.amazonaws.com/hrsl-cogs/hrsl_general/hrsl_general-latest.vrt"
TESELA_HRSL = "v1.5/cog_globallat_0_lon_-80_general-v1.5.4.tif"
URL_REGISTRO_HRSL = "https://registry.opendata.aws/dataforgood-fb-hrsl/"
URL_HDX_HRSL = "https://data.humdata.org/dataset/colombia-high-resolution-population-density-maps-demographic-estimates"
URL_DANE = "https://www.dane.gov.co/files/censo2018/proyecciones-de-poblacion/Municipal/PPED-AreaMun-2018-2042_VP.xlsx"
PAG_DANE = "https://www.dane.gov.co/index.php/estadisticas-por-tema/demografia-y-poblacion/proyecciones-de-poblacion"
URL_TERMINOS_DANE = ("https://www.dane.gov.co/index.php/servicios-al-ciudadano/tramites/"
                     "transparencia-y-acceso-a-la-informacion-publica/terminos-y-condiciones")
URL_MGN = "https://geoportal.dane.gov.co/mparcgis/rest/services/MGN2018/Serv_CapasMGN_2018/FeatureServer"
URL_MANZ = "https://geoportal.dane.gov.co/mparcgis/rest/services/MGN2018/Serv_CapaManzanaInt_2018/MapServer/0"
URL_SECR = "https://geoportal.dane.gov.co/mparcgis/rest/services/MGN2018/Serv_CapaSeccionRuralInt_2018/MapServer/0"
URL_INSTRUCTIVO = "https://geoportal.dane.gov.co/descargas/mgn-integrado/MGN2018_Integrado_CNPV2018_InstructivoUso.pdf"
URL_GEOVISOR_MANZ = "https://geoportal.dane.gov.co/geovisores/sociedad/cnpv2018-detallado/"
URL_GRILLA = "https://geoportal.dane.gov.co/mparcgis/rest/services/Grilla_DANE/Serv_Grilla_DANE/FeatureServer/1"
URL_FICHA_CNPV = "https://sitios.dane.gov.co/cnpv/app/views/informacion/fichas/76147.pdf"
URL_LEA = "https://conceptos.dane.gov.co/conceptos/conceptos/4720/ficha/"
URL_WP = "https://api.worldpop.org/v1/services/stats"
URL_WP_META = "https://hub.worldpop.org/geodata/summary?id=6349"
MPIO = "76147"           # DIVIPOLA de Cartago (Valle del Cauca)
ANIO_ANCLA = 2026        # año de la proyección DANE a la que se ancla la capa
ANIO_HRSL = 2020
ANIO_WP = 2020
SUBDIV = 5               # sub-rejilla de 0,00005° (≈ 5,5 m) para repartir las manzanas
UMBRAL_MANZANA_M2 = 20000.0   # manzanas ≥ 2 ha: reparto sobre edificaciones HRSL (supuesto documentado)
MODO_RURAL = "union"     # reparto rural dentro del cuadro de 1 km: unión de HRSL 2016 y WorldCover 2021 (ver repartir_rural)
OCUPADAS_MIN = 4         # umbral de anonimización del DANE (viviendas ocupadas con personas presentes)

CAMPOS_MANZ = ("COD_DANE_A,CLAS_CCDGO,SETR_CCNCT,SECR_CCNCT,SECU_CCNCT,AREA,TVIVIENDA,TP15_1_OCU,TP16_HOG,"
               "TP27_PERSO,PERSONAS_L,PERSONAS_S")
CAMPOS_SECR = "SECR_CCNCT,SETR_CCNCT,CLAS_CCDGO,AREA,STVIVIENDA,STP15_1_OC,TSP16_HOG,STP27_PERS,STPERSON_L,STPERSON_S"

CITA_HRSL = ("Meta and Center for International Earth Science Information Network - CIESIN - Columbia University. "
             "2022. High Resolution Settlement Layer (HRSL). Source imagery for HRSL © 2016 Maxar. "
             "Consultado en https://registry.opendata.aws/dataforgood-fb-hrsl el {hoy}.")
CITA_DANE = "Fuente: Departamento Administrativo Nacional de Estadística: www.dane.gov.co"
# ESA WorldCover 2021 v200: ventana local que guarda scripts/worldcover.py (no se descarga aquí)
WC_VENTANA = os.path.join(comun.FUENTES, "worldcover", "ESA_WorldCover_10m_2021_v200_N03W078_ventana_cartago.tif")
WC_CONSTRUIDO = 50       # clase «Built-up»: edificios, vías y otras estructuras (PUM v2.0, tabla 3)
URL_WC = "https://doi.org/10.5281/zenodo.7254221"
CITA_WC = ("Zanaga, D., Van De Kerchove, R., Daems, D., De Keersmaecker, W., Brockmann, C., Kirches, G., Wevers, J., "
           "Cartus, O., Santoro, M., Fritz, S., Lesiv, M., Herold, M., Tsendbazar, N.E., Xu, P., Ramoino, F., Arino, O., 2022. "
           "ESA WorldCover 10 m 2021 v200. doi:10.5281/zenodo.7254221")
ATRIB_WC = "© ESA WorldCover project 2021 / Contains modified Copernicus Sentinel data (2021) processed by ESA WorldCover consortium"
COBERTURA_HRSL_MIN = 0.5  # manzanas ≥ 2 ha: HRSL solo si cubre ≥ 50 % de lo construido según WorldCover (si no, unión)
CITA_WP = ("WorldPop (www.worldpop.org - School of Geography and Environmental Science, University of Southampton; "
           "Department of Geography and Geosciences, University of Louisville; Departement de Geographie, Universite de "
           "Namur) and Center for International Earth Science Information Network (CIESIN), Columbia University (2018). "
           "Global High Resolution Population Denominators Project - Funded by The Bill and Melinda Gates Foundation "
           "(OPP1134076). https://dx.doi.org/10.5258/SOTON/WP00645")
LIC_DANE = "Términos de uso del DANE: uso y transformación autorizados citando la fuente (" + URL_TERMINOS_DANE + ")"


def log(*a):
    print(*a, flush=True)


def n0(x):
    """Entero con punto de miles (convención en español)."""
    return f"{x:,.0f}".replace(",", ".")


def nd(x, d):
    """Número con d decimales, coma decimal y punto de miles."""
    return f"{x:,.{d}f}".replace(",", "§").replace(".", ",").replace("§", ".")


def reintentar(fn, *args, intentos=4, espera=3, **kw):
    ultimo = None
    for i in range(intentos):
        try:
            return fn(*args, **kw)
        except Exception as e:  # noqa: BLE001
            ultimo = e
            log(f"  {getattr(fn, '__name__', 'tarea')}: intento {i + 1}/{intentos} falló: {type(e).__name__}: {str(e)[:200]}")
            time.sleep(espera * (i + 1))
    raise ultimo


# ---------------------------------------------------------------- servicios ArcGIS del DANE

def _arcgis_query(url_capa, params):
    import requests
    base = {"outFields": "*", "returnGeometry": "true", "outSR": "4326", "f": "geojson"}
    base.update(params)
    # primera página sin parámetros de paginación (algunas capas del MGN, p. ej. la 305, los rechazan con HTTP 400);
    # solo si el servidor avisa que hay más registros se pide el resto con resultOffset
    feats, q = [], dict(base)
    while True:
        r = requests.get(url_capa + "/query", params=q, headers=UA, timeout=300)
        r.raise_for_status()
        j = r.json()
        if "error" in j:
            raise RuntimeError(j["error"])
        pagina = j.get("features", [])
        feats += pagina
        if not pagina or not (j.get("exceededTransferLimit") or j.get("properties", {}).get("exceededTransferLimit")):
            break
        q = dict(base, resultOffset=len(feats), resultRecordCount=len(pagina))
    return {"type": "FeatureCollection", "features": feats}


def arcgis(nombre, url_capa, params, refrescar=False):
    """Consulta paginada con caché. La caché guarda la consulta (URL y parámetros, incluida la caja) y la fecha; si la
    consulta pedida no coincide con la guardada (o la caché no la registra), se vuelve a consultar."""
    ruta = os.path.join(DIR, f"{nombre}.geojson")
    url_q = url_capa + "/query"
    if os.path.exists(ruta) and not refrescar:
        fc = json.load(open(ruta, encoding="utf-8"))
        c = fc.get("consulta") or {}
        if c.get("url") == url_q and c.get("parametros") == params:
            return fc
        log(f"{nombre}: la caché no registra esta misma consulta (URL, filtro o caja); se consulta de nuevo")
    log(f"DANE ArcGIS: consultando {nombre} …")
    fc = reintentar(_arcgis_query, url_capa, params)
    if not fc["features"]:
        raise RuntimeError(f"{nombre}: consulta sin resultados")
    fc["consulta"] = {"url": url_q, "parametros": params, "fecha": time.strftime("%Y-%m-%d"),
                      "registros": len(fc["features"])}
    comun.guardar_json(ruta, fc)
    return fc


def fecha_consulta(fc):
    return (fc or {}).get("consulta", {}).get("fecha")


def fecha_archivo(nombre):
    ruta = os.path.join(DIR, nombre)
    return time.strftime("%Y-%m-%d", time.localtime(os.path.getmtime(ruta))) if os.path.exists(ruta) else None


def mgn_dane(refrescar=False):
    from shapely.geometry import shape
    zu = arcgis("mgn2018_zona_urbana_76147", f"{URL_MGN}/305", {"where": f"COD_MPIO='{MPIO}'"}, refrescar)
    mu = arcgis("mgn2018_municipio_76147", f"{URL_MGN}/317",
                {"where": f"DPTO_CCDGO='{MPIO[:2]}' AND MPIO_CCDGO='{MPIO[2:]}'"}, refrescar)
    cab = [shape(f["geometry"]) for f in zu["features"] if f["properties"]["COD_CLAS"] == "1"]
    cps = [(f["properties"]["NOM_CPOB"], shape(f["geometry"]).buffer(0)) for f in zu["features"] if f["properties"]["COD_CLAS"] == "2"]
    assert len(cab) == 1 and len(mu["features"]) == 1, "MGN: se esperaba una cabecera y un municipio"
    return {"cabecera": cab[0].buffer(0), "municipio": shape(mu["features"][0]["geometry"]).buffer(0), "centros_poblados": cps,
            "fecha": {"zona_urbana": fecha_consulta(zu), "municipio": fecha_consulta(mu)}}


def grilla_dane(caja, refrescar=False):
    """Celdas de 1 km de la grilla DANE que tocan la ventana de lectura (rejilla ∪ municipio + margen)."""
    return arcgis("dane_grilla_1km_ventana", URL_GRILLA, {
        "geometry": ",".join(f"{x:.6f}" for x in caja), "geometryType": "esriGeometryEnvelope", "inSR": "4326",
        "spatialRel": "esriSpatialRelIntersects"}, refrescar)


def manzanas_dane(refrescar=False):
    return arcgis("cnpv2018_manzanas_76147", URL_MANZ,
                  {"where": f"MPIO_CDPMP='{MPIO}'", "outFields": CAMPOS_MANZ, "geometryPrecision": "7"}, refrescar)


def manzanas_vecinas_dane(caja, refrescar=False):
    """Manzanas censales de OTROS municipios dentro de la ventana (p. ej., Puerto Caldas, Pereira): sus personas se restan
    de la grilla de 1 km para que no se confundan con población rural de Cartago."""
    return arcgis("cnpv2018_manzanas_vecinas_ventana", URL_MANZ, {
        "where": f"MPIO_CDPMP<>'{MPIO}'", "geometry": ",".join(f"{x:.6f}" for x in caja), "geometryType": "esriGeometryEnvelope",
        "inSR": "4326", "spatialRel": "esriSpatialRelIntersects", "outFields": "COD_DANE_A,MPIO_CDPMP,CLAS_CCDGO,AREA,TP27_PERSO",
        "geometryPrecision": "7"}, refrescar)


def secciones_rurales_dane(refrescar=False):
    return arcgis("cnpv2018_secciones_rurales_76147", URL_SECR,
                  {"where": f"MPIO_CDPMP='{MPIO}'", "outFields": CAMPOS_SECR, "geometryPrecision": "7"}, refrescar)


def grilla_total_municipio(refrescar=False):
    """Total de la grilla de 1 km codificada como Cartago. Como arcgis(), la caché guarda la consulta (URL y parámetros) y
    solo se reutiliza si coincide con la pedida."""
    import requests
    ruta = os.path.join(DIR, "dane_grilla_1km_total_76147.json")
    url_q = URL_GRILLA + "/query"
    params = {"where": f"mpio_cod='{MPIO}'", "outFields": "personas,viviendas,hogares", "returnGeometry": "false", "f": "json"}
    if os.path.exists(ruta) and not refrescar:
        tot = json.load(open(ruta, encoding="utf-8"))
        c = tot.get("consulta") or {}
        if tot.get("fecha") and c.get("url") == url_q and c.get("parametros") == params:
            return tot
        log("dane_grilla_1km_total_76147: la caché no registra esta misma consulta (URL o filtro); se consulta de nuevo")
    def _q():
        r = requests.get(url_q, params=params, headers=UA, timeout=180)
        r.raise_for_status()
        j = r.json()
        if "error" in j:
            raise RuntimeError(j["error"])
        if j.get("exceededTransferLimit"):
            raise RuntimeError("la consulta excede el límite de registros del servicio: el total quedaría incompleto")
        fs = j["features"]
        if not fs:
            raise RuntimeError("consulta sin resultados")
        return {k: int(sum(f["attributes"][k] or 0 for f in fs)) for k in ("personas", "viviendas", "hogares")} | {"celdas": len(fs)}
    tot = reintentar(_q) | {"fecha": time.strftime("%Y-%m-%d"),
                            "consulta": {"url": url_q, "parametros": params, "fecha": time.strftime("%Y-%m-%d")}}
    comun.guardar_json(ruta, tot)
    return tot


# ---------------------------------------------------------------- HRSL

def caja_lectura(margen=0.015):   # ≥ diagonal de una celda de 1 km: toda celda que toque el municipio cabe entera
    from shapely.geometry import box
    mun_osm = comun._poligono(comun.osm()["municipio"])
    b = mun_osm.union(box(comun.OESTE, comun.SUR, comun.ESTE, comun.NORTE)).bounds
    return (b[0] - margen, b[1] - margen, b[2] + margen, b[3] + margen)


def _leer_hrsl_remoto(caja):
    import rasterio
    from rasterio.windows import from_bounds
    with rasterio.open("/vsicurl/" + URL_VRT) as src:
        assert src.crs.to_epsg() == 4326, src.crs
        w = from_bounds(*caja, transform=src.transform).round_offsets().round_lengths()
        arr = src.read(1, window=w)
        return arr, src.window_transform(w), {"res_grados": src.res[0], "nodata": str(src.nodata), "dtype": src.dtypes[0]}


def leer_hrsl(caja, refrescar=False):
    from affine import Affine
    ruta = os.path.join(DIR, "hrsl_ventana.npz")
    if os.path.exists(ruta) and not refrescar:
        z = np.load(ruta, allow_pickle=False)
        if np.allclose(z["caja"], caja):
            log(f"HRSL: ventana desde caché {ruta}")
            return z["arr"], Affine(*z["tr"][:6]), json.loads(str(z["info"]))
    log("HRSL: leyendo ventana remota del VRT …")
    t = time.time()
    arr, tr, info = reintentar(_leer_hrsl_remoto, caja)
    log(f"HRSL: {arr.shape} leída en {nd(time.time() - t, 1)} s")
    np.savez_compressed(ruta, arr=arr, tr=np.array(tuple(tr)), caja=np.array(caja), info=json.dumps(info))
    return arr, tr, info


def worldcover_construido(forma, transform):
    """Fracción construida (clase 50) de ESA WorldCover 2021 en una rejilla EPSG:4326 dada, por promedio de los píxeles de
    10 m; NaN donde la ventana local de WorldCover (rejilla común + 0,01°) no llega. None si falta la caché de
    scripts/worldcover.py (este script no la descarga)."""
    if not os.path.exists(WC_VENTANA):
        return None
    import rasterio
    from rasterio.warp import reproject, Resampling
    with rasterio.open(WC_VENTANA) as s:
        a = s.read(1)
        b = np.where(a == 0, np.nan, (a == WC_CONSTRUIDO).astype(np.float32)).astype(np.float32)
        dst = np.full(forma, np.nan, np.float32)
        reproject(source=b, destination=dst, src_transform=s.transform, src_crs=s.crs, src_nodata=np.nan,
                  dst_transform=transform, dst_crs="EPSG:4326", dst_nodata=np.nan, resampling=Resampling.average)
    return dst.astype(np.float64)


def pesos_solape(bordes_src, bordes_dst):
    """Matriz (n_dst, n_src): fracción de cada píxel fuente que cae en cada celda destino (1 dimensión)."""
    lo = np.maximum(bordes_dst[:, None, 0], bordes_src[None, :, 0])
    hi = np.minimum(bordes_dst[:, None, 1], bordes_src[None, :, 1])
    return np.clip(hi - lo, 0, None) / (bordes_src[:, 1] - bordes_src[:, 0])[None, :]


def agregar_conservando(p, tr):
    """Agrega personas/píxel (EPSG:4326) a la rejilla común repartiendo por área de solape."""
    ny, nx = p.shape
    xs = tr.c + np.arange(nx + 1) * tr.a
    ys = tr.f + np.arange(ny + 1) * tr.e          # tr.e < 0: de norte a sur
    dx = comun.OESTE + np.arange(comun.ANCHO + 1) * comun.RES
    dy = comun.NORTE - np.arange(comun.ALTO + 1) * comun.RES
    wx = pesos_solape(np.stack([xs[:-1], xs[1:]], 1), np.stack([dx[:-1], dx[1:]], 1))   # (ANCHO, nx)
    wy = pesos_solape(np.stack([ys[1:], ys[:-1]], 1), np.stack([dy[1:], dy[:-1]], 1))   # (ALTO, ny)
    salida = wy @ p @ wx.T
    dentro = (wy.sum(0)[:, None] * wx.sum(0)[None, :] * p).sum()   # masa fuente que cae en la rejilla
    return salida, float(dentro)


def agregar_rasterio_sum(p, tr):
    from rasterio.warp import reproject, Resampling
    dst = np.zeros((comun.ALTO, comun.ANCHO), dtype=np.float64)
    reproject(source=p.astype(np.float64), destination=dst, src_transform=tr, src_crs="EPSG:4326",
              dst_transform=comun.transformacion(), dst_crs="EPSG:4326", resampling=Resampling.sum)
    return dst


def mascara(geom, forma, transform):
    from rasterio.features import rasterize
    return rasterize([(geom, 1)], out_shape=forma, transform=transform, fill=0, dtype="uint8").astype(bool)


def indices_grilla(grilla, forma, transform):
    from rasterio.features import rasterize
    from shapely.geometry import shape
    return rasterize([(shape(f["geometry"]), k + 1) for k, f in enumerate(grilla["features"])], out_shape=forma,
                     transform=transform, fill=0, dtype="int32")


def redistribuir_en_grilla(p, idx, G):
    """Método de la versión anterior (solo para comparar): reparte G (personas por celda de 1 km) entre los píxeles
    HRSL de cada celda; si una celda tiene personas y ningún píxel HRSL, las reparte uniformemente."""
    n = len(G) + 1
    Gk = np.zeros(n)
    Gk[1:] = G
    H = np.bincount(idx.ravel(), weights=p.ravel(), minlength=n)
    N = np.bincount(idx.ravel(), minlength=n)
    Hi, Gi = H[idx], Gk[idx]
    q = np.where(Hi > 0, Gi * p / np.where(Hi > 0, Hi, 1.0), 0.0)
    unif = (idx > 0) & (Hi == 0) & (Gi > 0)
    q[unif] = (Gi / np.maximum(N[idx], 1))[unif]
    return q


# ---------------------------------------------------------------- censo 2018 por manzana y sección rural

def transform_sub():
    """Sub-rejilla SUBDIV × SUBDIV de la rejilla común (≈ 5,5 m) para repartir las manzanas."""
    from rasterio.transform import from_origin
    return (comun.ALTO * SUBDIV, comun.ANCHO * SUBDIV), from_origin(comun.OESTE, comun.NORTE, comun.RES / SUBDIV, comun.RES / SUBDIV)


def repartir_manzanas(geoms, P, area_m2, p, tr, umbral_m2=UMBRAL_MANZANA_M2, wc_sub=None):
    """Reparte las personas de cada manzana sobre su polígono en una sub-rejilla SUBDIV × SUBDIV y suma a la rejilla.

    Uniforme sobre el polígono, salvo en manzanas ≥ umbral_m2:
      - si las edificaciones HRSL (2016) cubren ≥ COBERTURA_HRSL_MIN del área construida según WorldCover 2021 dentro de la
        manzana (o WorldCover no ve nada construido), solo sobre las sub-celdas con edificación HRSL;
      - si no la cubren (HRSL desactualizado o incompleto), sobre la unión de sub-celdas HRSL y WorldCover construido;
      - sin ninguna de las dos señales, uniforme.
    Sin wc_sub (WorldCover ausente) se usa la regla anterior: HRSL si la manzana tiene alguna edificación HRSL.
    Manzanas sin ninguna sub-celda (muy pequeñas): todo a la celda de su punto interior.
    Devuelve (rejilla, info). La masa se conserva: Σ rejilla + fuera = Σ P.
    """
    from rasterio.features import rasterize
    S = SUBDIV
    (H, W), Tf = transform_sub()
    ids = rasterize([(g, k + 1) for k, g in enumerate(geoms)], out_shape=(H, W), transform=Tf, fill=0, dtype="int32")
    lat = comun.NORTE - (np.arange(H) + 0.5) * comun.RES / S
    lon = comun.OESTE + (np.arange(W) + 0.5) * comun.RES / S
    r = np.floor((lat - tr.f) / tr.e).astype(int)
    c = np.floor((lon - tr.c) / tr.a).astype(int)
    assert r.min() >= 0 and r.max() < p.shape[0] and c.min() >= 0 and c.max() < p.shape[1], "HRSL no cubre la rejilla"
    hb = p[np.ix_(r, c)] > 0
    n = len(geoms) + 1
    fl = ids.ravel()
    N = np.bincount(fl, minlength=n)
    NB = np.bincount(fl, weights=hb.ravel().astype(float), minlength=n)
    grande = np.zeros(n, bool)
    grande[1:] = np.asarray(area_m2) >= umbral_m2
    if wc_sub is not None:
        wb = np.nan_to_num(wc_sub) >= 0.5
        NW = np.bincount(fl, weights=wb.ravel().astype(float), minlength=n)
        NHW = np.bincount(fl, weights=(hb & wb).ravel().astype(float), minlength=n)
        cob = np.where(NW > 0, NHW / np.maximum(NW, 1), 1.0)
        usar_h = grande & (NB > 0) & (cob >= COBERTURA_HRSL_MIN)
        usar_u = grande & ~usar_h & ((NB + NW) > 0)
        w = np.where(usar_h[ids], hb, np.where(usar_u[ids], hb | wb, True)) & (ids > 0)
    else:
        cob = np.full(n, np.nan)
        usar_h = grande & (NB > 0)
        usar_u = np.zeros(n, bool)
        w = np.where(usar_h[ids], hb, True) & (ids > 0)
    Ws = np.bincount(fl, weights=w.ravel().astype(float), minlength=n)
    Pm = np.zeros(n)
    Pm[1:] = P
    val = np.where(w, Pm[ids] / np.where(Ws[ids] > 0, Ws[ids], 1.0), 0.0)
    rej = val.reshape(comun.ALTO, S, comun.ANCHO, S).sum((1, 3))
    sin, fuera = 0, 0.0
    for k in np.nonzero((N[1:] == 0) & (np.asarray(P) > 0))[0]:
        pt = geoms[k].representative_point()
        i, j = celda(pt.y, pt.x)
        if 0 <= i < comun.ALTO and 0 <= j < comun.ANCHO:
            rej[i, j] += P[k]
            sin += 1
        else:
            fuera += P[k]
    P = np.asarray(P, float)
    info = {"manzanas": int(len(geoms)), "personas": round(float(P.sum()), 1),
            "manzanas_sin_subcelda_asignadas_a_su_punto": int(sin),
            "personas_fuera_de_la_rejilla": round(fuera, 1),
            "manzanas_grandes": int(grande.sum()), "personas_en_manzanas_grandes": round(float(Pm[grande].sum()), 1),
            "manzanas_repartidas_sobre_edificaciones_hrsl": int(usar_h.sum()),
            "personas_en_esas_manzanas": round(float(Pm[usar_h].sum()), 1),
            "manzanas_repartidas_sobre_union_hrsl_worldcover": int(usar_u.sum()),
            "personas_en_manzanas_con_union": round(float(Pm[usar_u].sum()), 1),
            "manzanas_grandes_uniformes": int((grande & ~usar_h & ~usar_u).sum()),
            "manzanas_grandes_sin_edificacion_hrsl": int((grande & (NB == 0)).sum()),
            "conservacion_dif": round(float(rej.sum() + fuera - P.sum()), 6)}
    if wc_sub is not None:
        g_wc = grande & (NW > 0) & (NB > 0)
        info["cobertura_hrsl_de_lo_construido_worldcover_en_manzanas_grandes"] = {
            "umbral": COBERTURA_HRSL_MIN, "manzanas_con_ambas_senales": int(g_wc.sum()),
            "mediana": round(float(np.median(cob[g_wc])), 3) if g_wc.any() else None,
            "manzanas_bajo_el_umbral": int((g_wc & (cob < COBERTURA_HRSL_MIN)).sum()),
            "personas_en_manzanas_bajo_el_umbral": round(float(Pm[g_wc & (cob < COBERTURA_HRSL_MIN)].sum()), 1)}
    return rej, info


def personas_manzana_por_celda(geoms, P, celdas):
    """Personas de manzanas que caen en cada celda de 1 km (proporcional al área de la manzana dentro de la celda)."""
    from shapely.strtree import STRtree
    tree = STRtree(celdas)
    M = np.zeros(len(celdas))
    for g, pm in zip(geoms, P):
        if pm <= 0:
            continue
        a = g.area
        for j in tree.query(g):
            inter = g.intersection(celdas[j]).area
            if inter > 0:
                M[j] += pm * inter / a
    return M


def imputar_suprimidas(grilla, p_ocu, pers_por_ocu, kmax=OCUPADAS_MIN - 1):
    """Personas por celda de la grilla con las celdas suprimidas imputadas.

    Suprimida = personas 0 y viviendas > 0 (ninguna celda publicada tiene < 4 viviendas ocupadas). Imputación:
    E[ocupadas | ocupadas ≤ 3], ocupadas ~ Binomial(viviendas, p_ocu), × personas por vivienda ocupada.
    """
    from scipy.stats import binom
    G, sup, lo, hi = [], [], [], []
    for f in grilla["features"]:
        pr = f["properties"]
        pers, viv = int(pr.get("personas") or 0), int(pr.get("viviendas") or 0)
        if pers == 0 and viv > 0:
            k = np.arange(0, min(viv, kmax) + 1)
            pk = binom.pmf(k, viv, p_ocu)
            G.append(float((k * pk).sum() / pk.sum()) * pers_por_ocu)
            sup.append(True)
            lo.append(0.0)
            hi.append(min(viv, kmax) * pers_por_ocu)
        else:
            G.append(float(pers))
            sup.append(False)
            lo.append(float(pers))
            hi.append(float(pers))
    return np.array(G), np.array(sup), np.array(lo), np.array(hi)


def repartir_rural(p, idx, R, urb, sec, RD, c=None, modo="hrsl"):
    """Reparte la población rural dispersa de cada sección rural entre sus píxeles (resolución nativa de HRSL).

    R_k = personas rurales del cuadro k de 1 km (grilla con suprimidas imputadas − personas de TODAS las manzanas
    censales del cuadro, de Cartago y de municipios vecinos, con una tolerancia de ruido). Dentro de la parte no urbana
    del cuadro, R_k se reparte según un peso relativo u del píxel:
      modo «hrsl»        u = HRSL_píxel / Σ HRSL del cuadro (edificaciones de 2016);
      modo «worldcover»  u = construido_píxel / Σ construido del cuadro (ESA WorldCover 2021, fracción de clase 50);
      modo «union»       u = max(HRSL_píxel / Σ HRSL, construido_píxel / Σ construido): un píxel recibe población si
                         CUALQUIERA de las dos fuentes ve construcción (máximo de las dos participaciones normalizadas).
    WorldCover (c, fracción construida en la rejilla nativa) solo se usa en cuadros que la ventana de WorldCover cubre
    por completo; en los demás, HRSL. Si el cuadro no tiene ninguna señal, R_k se reparte de forma uniforme. Dentro de
    cada sección: RD_s × peso / Σ pesos de la sección, de modo que el total censal de cada sección se conserva exactamente.
    """
    assert modo in ("hrsl", "worldcover", "union"), modo
    n = len(R) + 1
    Rk = np.zeros(n)
    Rk[1:] = R
    nu = (~urb) & (idx > 0)
    Hk = np.bincount(idx[nu], weights=p[nu], minlength=n)
    Nk = np.bincount(idx[nu], minlength=n)
    hn = np.where(nu & (Hk[idx] > 0), p / np.where(Hk[idx] > 0, Hk[idx], 1.0), 0.0)
    cub = np.zeros(n, bool)
    Ck = np.zeros(n)
    if c is not None and modo != "hrsl":
        falta = np.bincount(idx[nu], weights=(~np.isfinite(c[nu])).astype(float), minlength=n)
        cub = (falta == 0) & (Nk > 0)
        cc = np.where(np.isfinite(c), c, 0.0)
        Ck = np.bincount(idx[nu], weights=cc[nu], minlength=n)
        cn = np.where(nu & cub[idx] & (Ck[idx] > 0), cc / np.where(Ck[idx] > 0, Ck[idx], 1.0), 0.0)
        if modo == "union":
            u = np.where(cub[idx], np.maximum(hn, cn), hn)
        else:
            u = np.where(cub[idx] & (Ck[idx] > 0), cn, hn)
    else:
        u = hn
    u = np.where(nu, u, 0.0)
    Uk = np.bincount(idx[nu], weights=u[nu], minlength=n)
    w = np.zeros(p.shape)
    con_u = nu & (Uk[idx] > 0)
    w[con_u] = (Rk[idx] * u / np.where(Uk[idx] > 0, Uk[idx], 1.0))[con_u]
    sin_u = nu & (Uk[idx] == 0) & (Rk[idx] > 0)
    w[sin_u] = (Rk[idx] / np.maximum(Nk[idx], 1))[sin_u]
    elig = (sec > 0) & ~urb
    out = np.zeros(p.shape)
    detalle, respaldo = [], []
    for s in range(1, len(RD) + 1):
        m = elig & (sec == s)
        Ws = float(w[m].sum())
        if Ws > 0:
            out[m] = RD[s - 1] * w[m] / Ws
        elif p[m].sum() > 0:
            out[m] = RD[s - 1] * p[m] / p[m].sum()
            respaldo.append(s)
        elif m.any():
            out[m] = RD[s - 1] / m.sum()
            respaldo.append(s)
        detalle.append({"seccion": s, "rd_censo_2018": round(float(RD[s - 1]), 1), "grilla_implicita": round(Ws, 1)})
    sh = ((Rk > 0) & (Hk == 0))[1:]
    su = ((Rk > 0) & (Uk == 0))[1:]
    cw = ((Rk > 0) & cub)[1:]
    return out, {"modo": modo, "secciones": detalle, "secciones_con_respaldo_hrsl_o_uniforme": respaldo,
                 "cuadros_sin_edificacion_hrsl": int(sh.sum()),
                 "personas_rurales_en_cuadros_sin_edificacion_hrsl": round(float(Rk[1:][sh].sum()), 1),
                 "cuadros_sin_edificacion_hrsl_reparto_uniforme": int((sh & su).sum()),   # nombre de versiones anteriores
                 "personas_rurales_en_esos_cuadros": round(float(Rk[1:][sh & su].sum()), 1),
                 "cuadros_sin_ninguna_senal_reparto_uniforme": int(su.sum()),
                 "personas_rurales_en_cuadros_uniformes": round(float(Rk[1:][su].sum()), 1),
                 "cuadros_con_worldcover": int(cw.sum()),
                 "personas_rurales_en_cuadros_con_worldcover": round(float(Rk[1:][cw].sum()), 1),
                 "_uniformes": su, "_hk": Hk[1:], "_uk": Uk[1:]}


# ---------------------------------------------------------------- DANE (Excel y ficha)

def _descargar(url, ruta):
    import requests
    tmp = ruta + ".parcial"
    hecho = os.path.getsize(tmp) if os.path.exists(tmp) else 0
    cab = dict(UA, Range=f"bytes={hecho}-") if hecho else UA
    with requests.get(url, headers=cab, stream=True, timeout=120) as r:
        r.raise_for_status()
        total = int(r.headers.get("Content-Length", 0)) + (hecho if r.status_code == 206 else 0)
        with open(tmp, "ab" if r.status_code == 206 else "wb") as f:
            for trozo in r.iter_content(1 << 16):
                f.write(trozo)
    if total and os.path.getsize(tmp) < total:
        raise IOError(f"descarga incompleta {os.path.getsize(tmp)}/{total} bytes")   # el reintento continúa con Range
    os.replace(tmp, ruta)


def leer_dane(refrescar=False):
    import openpyxl   # instalado para este script (uv pip install openpyxl): lectura del .xlsx del DANE
    ruta = os.path.join(DIR, "PPED-AreaMun-2018-2042_VP.xlsx")
    if refrescar or not os.path.exists(ruta):
        log("DANE: descargando Excel de proyecciones municipales por área (~3,9 MB) …")
        reintentar(_descargar, URL_DANE, ruta, intentos=6)
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    actualizado = next((f[0].strip() for f in wb["PPED"].iter_rows(values_only=True)
                        if f and isinstance(f[0], str) and f[0].startswith("Actualizado")), None)
    serie, encabezado = {}, None
    clases = {"Cabecera Municipal": "cabecera", "Centros Poblados y Rural Disperso": "resto", "Total": "total"}
    for fila in wb["PobMunicipalxÁrea"].iter_rows(values_only=True):
        if fila and fila[0] == "DP":
            encabezado = list(fila)
        elif encabezado and fila and fila[2] == MPIO:
            serie.setdefault(int(fila[4]), {})[clases[fila[5]]] = int(fila[6])
    assert encabezado == ["DP", "DPNOM", "MPIO", "DPMP", "AÑO", "ÁREA GEOGRÁFICA", "TOTAL"], encabezado
    assert serie, "Cartago (76147) no aparece en el Excel del DANE"
    for anio, v in serie.items():
        assert v["cabecera"] + v["resto"] == v["total"], (anio, v)
    return {"serie": dict(sorted(serie.items())), "actualizado": actualizado}


def ficha_cnpv(refrescar=False):
    """Personas censadas CNPV 2018 por clase, en hogares particulares y en LEA (ficha municipal). Requiere pdftotext."""
    ruta = os.path.join(DIR, "cnpv2018_ficha_76147.pdf")
    if refrescar or not os.path.exists(ruta):
        reintentar(_descargar, URL_FICHA_CNPV, ruta)
    if not shutil.which("pdftotext"):
        return None
    txt = subprocess.run(["pdftotext", "-layout", ruta, "-"], capture_output=True, text=True, check=True).stdout
    num = r"\s+([\d.]+)"
    tot = re.search(r"^\s*Cartago" + num * 6 + r"\s*$", txt, re.M)
    cab = re.search(r"Cabecera" + num * 6 + r"[^\n]*\n\s*Cartago\s+Centro poblado y\s*\n" + num * 6, txt)
    if not (tot and cab):
        return None
    n = lambda s: int(s.replace(".", ""))  # noqa: E731
    return {"total_censado_2018": n(tot.group(1)), "total_censado_2005": n(tot.group(2)),
            "hogares_particulares_2018": n(tot.group(3)), "lea_2018": n(tot.group(5)),
            "cabecera_2018": n(cab.group(1)), "cabecera_hp_2018": n(cab.group(3)), "cabecera_hogares_2018": n(cab.group(5)),
            "resto_2018": n(cab.group(7)), "resto_hp_2018": n(cab.group(9)), "resto_hogares_2018": n(cab.group(11))}


# ---------------------------------------------------------------- WorldPop

def _geojson(geom):
    from shapely.geometry import mapping
    return json.dumps({"type": "FeatureCollection", "features": [
        {"type": "Feature", "properties": {}, "geometry": mapping(geom)}]}, separators=(",", ":"))


def _worldpop_una(geojson):
    import requests
    # POST: con GET, el GeoJSON del municipio excede el largo de URL admitido (HTTP 414)
    r = requests.post(URL_WP, data={"dataset": "wpgppop", "year": ANIO_WP, "geojson": geojson}, timeout=120)
    r.raise_for_status()
    j = r.json()
    if j.get("error"):
        raise RuntimeError(j.get("error_message"))
    if j.get("status") != "finished":
        for _ in range(60):
            time.sleep(3)
            j = requests.get(f"https://api.worldpop.org/v1/tasks/{j['taskid']}", timeout=60).json()
            if j.get("status") in ("finished", "failed"):
                break
    if j.get("status") != "finished" or j.get("error"):
        raise RuntimeError(f"WorldPop: {j.get('status')} {j.get('error_message')}")
    return float(j["data"]["total_population"])


def worldpop(geoms, refrescar=False):
    """geoms: {nombre: shapely} → {nombre: total o None}. Caché por hash del GeoJSON."""
    ruta = os.path.join(DIR, "worldpop.json")
    cache = {} if refrescar or not os.path.exists(ruta) else json.load(open(ruta, encoding="utf-8"))
    claves = {n: f"{n}:{ANIO_WP}:{hashlib.sha1(_geojson(g).encode()).hexdigest()[:12]}" for n, g in geoms.items()}
    pend = {c: _geojson(geoms[n]) for n, c in claves.items() if c not in cache}
    if pend:
        log(f"WorldPop: consultando {len(pend)} polígono(s) …")
        with ThreadPoolExecutor(3) as ex:
            fut = {c: ex.submit(reintentar, _worldpop_una, gj) for c, gj in pend.items()}
            for c, f in fut.items():
                try:
                    cache[c] = f.result()
                except Exception as e:  # noqa: BLE001
                    log(f"WorldPop: falló {c}: {e}")
        comun.guardar_json(ruta, cache)
    return {n: cache.get(c) for n, c in claves.items()}


# ---------------------------------------------------------------- utilidades de verificación

def celda(lat, lon):
    return int((comun.NORTE - lat) / comun.RES), int((lon - comun.OESTE) / comun.RES)


def decodificar_png(nombre, escala, desplazamiento):
    from PIL import Image
    im = np.asarray(Image.open(os.path.join(comun.CAPAS, f"{nombre}.png")).convert("RGBA")).astype(np.uint32)
    v = im[..., 0] * 256 + im[..., 1]
    return np.where(v == 0, np.nan, (v.astype(np.float64) - 1) * escala + desplazamiento)


def muestrear_linea(arr, linea, paso_m=20.0):
    n = max(2, int(linea.length * 111000 / paso_m))
    out = []
    for k in range(n):
        pt = linea.interpolate(k / (n - 1), normalized=True)
        i, j = celda(pt.y, pt.x)
        if 0 <= i < comun.ALTO and 0 <= j < comun.ANCHO:
            out.append(arr[i, j])
    return np.array(out, dtype=float)


def indice_disimilitud(a, b):
    """½ Σ |a_i/Σa − b_i/Σb|: fracción de la población que habría que mover para igualar los repartos."""
    return float(0.5 * np.abs(a / a.sum() - b / b.sum()).sum())


def bloques(a, b):
    H, W = a.shape
    Hp, Wp = -(-H // b) * b, -(-W // b) * b
    z = np.zeros((Hp, Wp))
    z[:H, :W] = np.nan_to_num(a)
    return z.reshape(Hp // b, b, Wp // b, b).sum((1, 3))


def comparar_por_bloques(T, X, m, b, minimo=100):
    """Compara el reparto X con la referencia T dentro de la máscara m, en bloques de b × b celdas."""
    from scipy.stats import pearsonr
    t, x = bloques(np.where(m, T, 0), b), bloques(np.where(m, X, 0), b)
    k = (t > 0) | (x > 0)
    t, x = t[k], x[k]
    xs = x * t.sum() / x.sum()
    g = t >= minimo
    err = np.abs(xs[g] / t[g] - 1)
    return {"bloques": int(k.sum()), "pearson": round(float(pearsonr(t, x)[0]), 3),
            "indice_disimilitud": round(indice_disimilitud(t, x), 3),
            f"bloques_con_ref_{minimo}_o_mas": int(g.sum()),
            "error_relativo_mediano": round(float(np.median(err)), 3),
            "frac_bloques_error_mayor_50pct": round(float((err > 0.5).mean()), 3)}


def anillo(geom, d0_m, d1_m, forma, transform):
    g1 = geom.buffer(d1_m / 111000)
    m = mascara(g1, forma, transform)
    if d0_m > 0:
        m &= ~mascara(geom.buffer(d0_m / 111000), forma, transform)
    return m


# ---------------------------------------------------------------- principal

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--refrescar", action="store_true", help="ignora las cachés y vuelve a descargar")
    args = ap.parse_args()
    hoy = time.strftime("%Y-%m-%d")
    from scipy.stats import pearsonr
    from shapely.geometry import box, shape, mapping, LineString
    from pyproj import Transformer
    from shapely.ops import transform as transformar

    utm = Transformer.from_crs("EPSG:4326", "EPSG:32618", always_xy=True).transform
    km2 = lambda g: transformar(utm, g).area / 1e6  # noqa: E731
    rejilla = box(comun.OESTE, comun.SUR, comun.ESTE, comun.NORTE)
    caja = caja_lectura()
    FORMA = (comun.ALTO, comun.ANCHO)
    T = comun.transformacion()

    # 1. lecturas remotas en paralelo
    def opcional(f, nombre):
        try:
            return f.result()
        except Exception as e:  # noqa: BLE001
            log(f"{nombre} no disponible: {type(e).__name__}: {str(e)[:300]}")
            return None

    with ThreadPoolExecutor(6) as ex:
        f_hrsl = ex.submit(leer_hrsl, caja, args.refrescar)
        f_dane = ex.submit(leer_dane, args.refrescar)
        f_mgn = ex.submit(mgn_dane, args.refrescar)
        f_man = ex.submit(manzanas_dane, args.refrescar)
        f_sec = ex.submit(secciones_rurales_dane, args.refrescar)
        f_vec = ex.submit(manzanas_vecinas_dane, caja, args.refrescar)
        f_gri = ex.submit(grilla_dane, caja, args.refrescar)
        f_gtot = ex.submit(grilla_total_municipio, args.refrescar)
        f_ficha = ex.submit(ficha_cnpv, args.refrescar)
        mgn = opcional(f_mgn, "MGN 2018 del DANE")
        if mgn:
            f_wp = ex.submit(worldpop, {"cabecera": mgn["cabecera"], "municipio": mgn["municipio"], "rejilla": rejilla},
                             args.refrescar)
        arr, tr, info_hrsl = f_hrsl.result()          # sin HRSL no hay capa: el error se propaga
        dane = opcional(f_dane, "Proyecciones DANE")
        manz = opcional(f_man, "Manzanas CNPV 2018")
        secr = opcional(f_sec, "Secciones rurales CNPV 2018")
        vecinas = opcional(f_vec, "Manzanas de municipios vecinos")
        grilla = opcional(f_gri, "Grilla DANE 1 km")
        grilla_tot = opcional(f_gtot, "Total grilla DANE")
        ficha = opcional(f_ficha, "Ficha CNPV 2018")
        wp = (opcional(f_wp, "WorldPop") or {}) if mgn else {}

    faltan = [n for n, v in (("MGN 2018", mgn), ("proyecciones PPED", dane), ("manzanas CNPV 2018", manz),
                             ("secciones rurales CNPV 2018", secr)) if not v]
    if faltan or ANIO_ANCLA not in dane["serie"]:
        raise SystemExit(f"Faltan fuentes oficiales ({', '.join(faltan) or f'año {ANIO_ANCLA}'}): no se publica la capa.")
    cab, mun = mgn["cabecera"], mgn["municipio"]
    d26 = dane["serie"][ANIO_ANCLA]

    b = mun.bounds
    assert caja[0] <= b[0] and caja[1] <= b[1] and caja[2] >= b[2] and caja[3] >= b[3], "la ventana HRSL no cubre el municipio"
    p = np.nan_to_num(arr.astype(np.float64), nan=0.0)
    assert np.all(p >= 0), "HRSL con valores negativos"

    # 2. HRSL: agregación conservando la masa (versión auxiliar y contraste)
    hrsl, masa_dentro = agregar_conservando(p, tr)
    suma_rio = float(agregar_rasterio_sum(p, tr).sum())
    yc = tr.f + (np.arange(p.shape[0]) + 0.5) * tr.e
    xc = tr.c + (np.arange(p.shape[1]) + 0.5) * tr.a
    dentro_c = ((yc > comun.SUR) & (yc < comun.NORTE))[:, None] & ((xc > comun.OESTE) & (xc < comun.ESTE))[None, :]
    masa_centros = float(p[dentro_c].sum())
    suma_rejilla = float(hrsl.sum())
    conserv = {
        "suma_rejilla_hab": round(suma_rejilla, 1),
        "masa_fuente_dentro_de_la_rejilla_hab": round(masa_dentro, 1),
        "masa_fuente_por_centros_hab": round(masa_centros, 1),
        "suma_rasterio_Resampling_sum_hab": round(suma_rio, 1),
        "dif_vs_masa_fuente_pct": round(100 * (suma_rejilla / masa_dentro - 1), 6) + 0.0,
        "dif_vs_centros_pct": round(100 * (suma_rejilla / masa_centros - 1), 4),
        "dif_rasterio_sum_vs_exacto_pct": round(100 * (suma_rio / suma_rejilla - 1), 4),
    }
    log("Conservación HRSL:", conserv)
    if max(abs(conserv["dif_vs_centros_pct"]), abs(conserv["dif_vs_masa_fuente_pct"])) > 2:
        raise SystemExit("La agregación no conserva la población dentro de ±2 %: no se publica la capa.")
    np.save(os.path.join(comun.REJILLA_NPY, "poblacion_hrsl.npy"), hrsl.astype(np.float32))

    m_mun_nat = mascara(mun, p.shape, tr)
    m_cab_nat = mascara(cab, p.shape, tr)
    h_mun = float(p[m_mun_nat].sum())
    h_cab_nat = float(p[m_cab_nat].sum())
    h_mun_fuera = float(p[m_mun_nat & ~mascara(rejilla, p.shape, tr)].sum())
    m_cab = mascara(cab, FORMA, T)
    m_mun = mascara(mun, FORMA, T) | m_cab
    m_resto = m_mun & ~m_cab

    # 3a. manzanas CNPV 2018
    fm = manz["features"]
    g_man = [shape(f["geometry"]).buffer(0) for f in fm]
    pr_man = [f["properties"] for f in fm]
    # Instructivo MGN2018 integrado (tabla de manzana): TP27_PERSO = «Número de personas», PERSONAS_L = «Conteo de personas
    # en LEAS», PERSONAS_S = «Conteo de personas en hogares particulares». Se usa PERSONAS_S; los LEA se publican solo por
    # clase, así que PERSONAS_L = 0 y TP27_PERSO = PERSONAS_S en cada manzana (se comprueba).
    P_man = np.array([float(x["PERSONAS_S"] or 0) for x in pr_man])
    P_tot_man = np.array([float(x["TP27_PERSO"] or 0) for x in pr_man])
    cls_man = np.array([x["CLAS_CCDGO"] for x in pr_man])
    area_man = np.array([float(x["AREA"] or 0) for x in pr_man])
    lea_man = float(sum(float(x.get("PERSONAS_L") or 0) for x in pr_man))
    assert np.array_equal(P_tot_man, P_man + np.array([float(x.get("PERSONAS_L") or 0) for x in pr_man])), \
        "manzanas: TP27_PERSO ≠ PERSONAS_S + PERSONAS_L"
    assert lea_man == 0 and np.array_equal(P_tot_man, P_man), "manzanas con personas en LEA: revisar el uso de PERSONAS_S"
    assert set(cls_man) <= {"1", "2"}, set(cls_man)
    k1, k2 = cls_man == "1", cls_man == "2"
    P1, P2 = float(P_man[k1].sum()), float(P_man[k2].sum())
    sec_feats = secr["features"]
    g_sec = [shape(f["geometry"]).buffer(0) for f in sec_feats]
    cod_sec = [f["properties"]["SECR_CCNCT"] for f in sec_feats]
    # secciones rurales: STP27_PERS = «Número de personas», STPERSON_S = «en hogares particulares», STPERSON_L = «en LEAS»
    T_sec = np.array([float(f["properties"]["STPERSON_S"] or 0) for f in sec_feats])
    T_tot_sec = np.array([float(f["properties"]["STP27_PERS"] or 0) for f in sec_feats])
    lea_sec = float(sum(float(f["properties"].get("STPERSON_L") or 0) for f in sec_feats))
    assert lea_sec == 0 and np.array_equal(T_tot_sec, T_sec), "secciones con personas en LEA: revisar el uso de STPERSON_S"
    viv_sec = float(sum(f["properties"]["STVIVIENDA"] or 0 for f in sec_feats))
    ocu_sec = float(sum(f["properties"]["STP15_1_OC"] or 0 for f in sec_feats))
    # manzanas de centro poblado → sección rural del mismo sector y sección: el código de la manzana lleva la clase 2
    # en la 6.ª posición (76147 2 002 02) y el de la sección, la clase 3 (76147 3 002 02). Comprobado con la geometría:
    # cada centro poblado queda dentro (o a < 110 m) de la sección así emparejada. Respaldo: la sección más cercana.
    C_sec = np.zeros(len(sec_feats))
    asign_cp = {"por_codigo": 0, "por_ubicacion": 0}
    for k in np.nonzero(k2)[0]:
        c = pr_man[k].get("SECR_CCNCT") or ""
        c3 = c[:5] + "3" + c[6:] if len(c) == 11 else c
        s = cod_sec.index(c3) if c3 in cod_sec else None
        if s is None:
            pt = g_man[k].representative_point()
            s = int(np.argmin([g.distance(pt) for g in g_sec]))
            asign_cp["por_ubicacion"] += 1
        else:
            asign_cp["por_codigo"] += 1
        C_sec[s] += P_man[k]
    RD = T_sec - C_sec
    assert np.all(RD >= 0), f"secciones con menos personas que sus centros poblados: {RD}"
    totales_censo = {"cabecera_manzanas": P1, "centros_poblados_manzanas": P2, "secciones_rurales": float(T_sec.sum()),
                     "rural_disperso": float(RD.sum()), "hogares_particulares": P1 + float(T_sec.sum()),
                     "lea_en_manzanas": lea_man, "lea_en_secciones": lea_sec, "asignacion_cp_a_secciones": asign_cp,
                     "campos": {"manzana": "PERSONAS_S (personas en hogares particulares); TP27_PERSO (número de personas) = "
                                           "PERSONAS_S en todas las manzanas porque PERSONAS_L (LEA) = 0",
                                "seccion_rural": "STPERSON_S; STP27_PERS = STPERSON_S porque STPERSON_L = 0",
                                "manzanas_vecinas": "TP27_PERSO (número de personas; los LEA no se publican por manzana)",
                                "fuente_de_las_definiciones": URL_INSTRUCTIVO}}
    if ficha:
        totales_censo["ficha"] = ficha
        totales_censo["coincide_cabecera_con_ficha"] = P1 == ficha["cabecera_hp_2018"]
        totales_censo["coincide_resto_con_ficha"] = float(T_sec.sum()) == ficha["resto_hp_2018"]
        totales_censo["coincide_lea_con_ficha"] = ficha["total_censado_2018"] - ficha["hogares_particulares_2018"] == ficha["lea_2018"]
    log("Totales censales 2018:", json.dumps(totales_censo, ensure_ascii=False))

    # tamaño de las secciones urbanas (la anonimización une manzanas dentro de su sección): diagonal del rectángulo
    from collections import defaultdict
    from pyproj import Geod
    geod = Geod(ellps="WGS84")
    por_secu = defaultdict(list)
    for g, x, k in zip(g_man, pr_man, k1):
        if k:
            por_secu[x.get("SECU_CCNCT")].append(g)
    diag = []
    for gs_ in por_secu.values():
        bx = [min(g.bounds[0] for g in gs_), min(g.bounds[1] for g in gs_), max(g.bounds[2] for g in gs_), max(g.bounds[3] for g in gs_)]
        diag.append(geod.inv(bx[0], bx[1], bx[2], bx[3])[2])
    vacias = int(sum(1 for x, k in zip(pr_man, k1) if k and (x["PERSONAS_S"] or 0) == 0 and (x["TVIVIENDA"] or 0) > 0))
    totales_censo["secciones_urbanas"] = {"n": len(diag), "diagonal_mediana_m": round(float(np.median(diag))),
                                          "manzanas_cabecera_con_viviendas_y_0_personas": vacias}
    # WorldCover 2021 (caché de scripts/worldcover.py): construido en la sub-rejilla de las manzanas y en la rejilla nativa
    # de HRSL para el reparto rural
    forma_sub, tr_sub = transform_sub()
    wc_sub = worldcover_construido(forma_sub, tr_sub)
    if wc_sub is None:
        log(f"WorldCover: falta {WC_VENTANA} (scripts/worldcover.py); manzanas grandes y rural solo con HRSL")
    g_k1, g_k2 = [g for g, k in zip(g_man, k1) if k], [g for g, k in zip(g_man, k2) if k]
    rej_c1, info_m1 = repartir_manzanas(g_k1, P_man[k1], area_man[k1], p, tr, wc_sub=wc_sub)
    rej_c2, info_m2 = repartir_manzanas(g_k2, P_man[k2], area_man[k2], p, tr, wc_sub=wc_sub)
    rej_c1_unif, _ = repartir_manzanas(g_k1, P_man[k1], area_man[k1], p, tr, umbral_m2=np.inf)
    rej_c1_solo_hrsl, _ = repartir_manzanas(g_k1, P_man[k1], area_man[k1], p, tr)   # regla anterior (solo HRSL)
    del wc_sub
    log("Manzanas cabecera:", info_m1)
    log("Manzanas centros poblados:", info_m2)
    assert abs(info_m1["conservacion_dif"]) < 1e-6 and abs(info_m2["conservacion_dif"]) < 1e-6
    assert info_m1["personas_fuera_de_la_rejilla"] == 0, "hay manzanas de la cabecera fuera de la rejilla"

    # 3b. población rural dispersa dentro de cada sección rural
    from rasterio.features import rasterize
    urb_geoms = [cab] + [g for _, g in mgn["centros_poblados"]]
    g_vec = [shape(f["geometry"]).buffer(0) for f in (vecinas or {}).get("features", [])]
    P_vec = np.array([float(f["properties"].get("TP27_PERSO") or 0) for f in (vecinas or {}).get("features", [])])
    urb_nat = rasterize([(g, 1) for g in urb_geoms + g_vec], out_shape=p.shape, transform=tr, fill=0, dtype="uint8").astype(bool)
    sec_nat = rasterize([(g, k + 1) for k, g in enumerate(g_sec)], out_shape=p.shape, transform=tr, fill=0, dtype="int32")
    p_ocu, ppo = ocu_sec / viv_sec, float(T_sec.sum()) / ocu_sec
    info_rural = {"tasa_ocupacion_rural": round(p_ocu, 4), "personas_por_vivienda_ocupada_rural": round(ppo, 3)}
    info_grilla, q_old_nat = None, None
    c_nat = worldcover_construido(p.shape, tr)
    modo_rural = MODO_RURAL if c_nat is not None else "hrsl"
    rd_alt = {}
    if grilla:
        gf = grilla["features"]
        celdas = [shape(f["geometry"]) for f in gf]
        G_raw = np.array([float(f["properties"].get("personas") or 0) for f in gf])
        G_imp, sup, G_lo, G_hi = imputar_suprimidas(grilla, p_ocu, ppo)
        M_cel = personas_manzana_por_celda(g_man + g_vec, np.r_[P_man, P_vec], celdas)
        idx_nat = indices_grilla(grilla, p.shape, tr)
        from shapely.ops import unary_union
        U = unary_union(urb_geoms)
        cel_mun = np.array([c.intersection(mun).area / c.area for c in celdas])
        cel_urb = np.array([c.intersection(U).area / c.area for c in celdas])
        # tolerancia de ruido: diferencia relativa máxima grilla vs manzanas en cuadros ≥ 95 % urbanos (mismo censo,
        # distinta agregación); por debajo de ella, el exceso de la grilla sobre las manzanas no se toma como población rural
        urbanos = (cel_urb >= 0.95) & (M_cel > 0)
        tau = float(np.max(np.abs(G_imp[urbanos] / M_cel[urbanos] - 1))) if urbanos.sum() >= 3 else 0.0
        R = np.clip(G_imp - (1 + tau) * M_cel, 0, None)
        rr = lambda Rx, modo=modo_rural: repartir_rural(p, idx_nat, Rx, urb_nat, sec_nat, RD, c=c_nat, modo=modo)  # noqa: E731
        rd_nat, info_rep = rr(R)
        rd_sin_tope, info_sin_tope = rr(np.clip(G_imp - M_cel, 0, None))
        # sensibilidad de la imputación: sin imputar (suprimidas = 0) y con el máximo (3 viviendas ocupadas)
        rd_lo, info_lo = rr(np.clip(G_lo - (1 + tau) * M_cel, 0, None))
        rd_hi, info_hi = rr(np.clip(G_hi - (1 + tau) * M_cel, 0, None))
        # imputación con razones de la PROPIA grilla: celdas rurales publicadas (degurba_l1 = 1; en la codificación GHSL del
        # grado de urbanización, 1 = celdas rurales y 3 = centro urbano; supuesto coherente con los datos: las celdas 3
        # tienen miles de personas por km²). Con ellas el contraste grilla-censo no usa los totales censales de las secciones.
        rur_pub = [f["properties"] for f in gf if str(f["properties"].get("degurba_l1")) == "1" and (f["properties"].get("personas") or 0) > 0]
        ppo_g = sum(x["personas"] for x in rur_pub) / sum(x["vivienda_o"] for x in rur_pub)
        p_ocu_g = sum(x["vivienda_o"] for x in rur_pub) / sum(x["viviendas"] for x in rur_pub)
        G_imp_g = imputar_suprimidas(grilla, p_ocu_g, ppo_g)[0]
        rd_g, info_g = rr(np.clip(G_imp_g - (1 + tau) * M_cel, 0, None))
        # alternativas de reparto dentro del cuadro (sensibilidad): solo HRSL 2016 y solo WorldCover 2021
        for modo in ("hrsl", "worldcover", "union"):
            if modo != modo_rural and (modo == "hrsl" or c_nat is not None):
                rd_alt[modo], _ = rr(R, modo)
        info_rural["cuadros_que_tocan_cartago_con_reparto_uniforme"] = int((info_rep["_uniformes"] & (cel_mun > 0)).sum())
        info_rural["cuadros_que_tocan_cartago_sin_edificacion_hrsl"] = int(((R > 0) & (info_rep["_hk"] == 0) & (cel_mun > 0)).sum())
        imp = lambda inf: round(sum(d["grilla_implicita"] for d in inf["secciones"]), 1)  # noqa: E731
        info_rural.update({"manzanas_vecinas_en_la_ventana": len(g_vec), "personas_manzanas_vecinas": round(float(P_vec.sum())),
                           "cuadros_urbanos_para_tolerancia": int(urbanos.sum()), "tolerancia_ruido_tau": round(tau, 4),
                           "rural_implicito_sin_tolerancia": imp(info_sin_tope), "rural_implicito_sin_imputar": imp(info_lo),
                           "rural_implicito_imputacion_maxima": imp(info_hi),
                           "imputacion_con_razones_de_la_grilla": {
                               "celdas": "celdas rurales publicadas de la ventana (degurba_l1 = 1, personas > 0)",
                               "n_celdas": len(rur_pub), "tasa_ocupacion": round(p_ocu_g, 4),
                               "personas_por_vivienda_ocupada": round(ppo_g, 3),
                               "personas_imputadas_en_celdas_que_tocan_cartago_ponderadas_por_area":
                                   round(float((G_imp_g * sup * cel_mun).sum()), 1),
                               "rural_implicito_en_secciones": imp(info_g),
                               "nota": "Razones calculadas solo con la grilla publicada, sin los totales censales de las secciones; "
                                       "las celdas publicadas tienen ≥ 4 viviendas ocupadas, así que pueden no representar a "
                                       "las suprimidas."}})
        for inf in (info_rep, info_sin_tope, info_lo, info_hi, info_g):
            for k_ in [k_ for k_ in inf if k_.startswith("_")]:
                del inf[k_]
        # parte del municipio de cada celda (para diagnóstico) y celdas suprimidas que tocan Cartago
        toca = cel_mun > 0
        mpio_c = np.array([f["properties"].get("mpio_cod") for f in gf])
        # grilla repartida con HRSL (método anterior): cuánto cae en el polígono municipal y de qué celdas viene
        q_old_nat = redistribuir_en_grilla(p, idx_nat, G_raw)
        n = len(gf) + 1
        por_celda_en_mun = np.bincount(idx_nat[m_mun_nat], weights=q_old_nat[m_mun_nat], minlength=n)[1:]
        info_grilla = {
            "celdas_en_la_ventana": len(gf), "fecha_consulta": fecha_consulta(grilla),
            "personas_publicadas_en_la_ventana": int(G_raw.sum()),
            "celdas_suprimidas_en_la_ventana": int(sup.sum()),
            "celdas_suprimidas_que_tocan_cartago": int((sup & toca).sum()),
            "personas_imputadas_en_celdas_que_tocan_cartago_ponderadas_por_area": round(float((G_imp * sup * cel_mun).sum()), 1),
            "rango_imputacion_ponderado_por_area": [0.0, round(float((G_hi * sup * cel_mun).sum()), 1)],
            "grilla_repartida_con_hrsl_dentro_del_municipio": round(float(por_celda_en_mun.sum()), 1),
            "de_celdas_codificadas_cartago": round(float(por_celda_en_mun[mpio_c == MPIO].sum()), 1),
            "de_celdas_de_otros_municipios": {m: round(float(por_celda_en_mun[mpio_c == m].sum()), 1)
                                              for m in sorted(set(mpio_c[(mpio_c != MPIO) & (por_celda_en_mun > 1)]))},
            "personas_manzanas_asignadas_a_celdas": round(float(M_cel.sum()), 1),
            "rural_implicito_grilla_menos_manzanas_en_secciones": imp(info_rep),
            "parametros_rurales": info_rural,
            "grilla_en_municipio_mas_imputacion": round(float(por_celda_en_mun.sum() + (G_imp * sup * cel_mun).sum()), 1),
            "hogares_particulares_cnpv2018": P1 + float(T_sec.sum()),
            "rural_disperso_censo_2018": round(float(RD.sum()), 1),
            "reparto": info_rep,
        }
        log("Grilla y rural:", json.dumps(info_grilla, ensure_ascii=False))
    else:
        elig = (sec_nat > 0) & ~urb_nat
        rd_nat = np.zeros(p.shape)
        for s in range(1, len(RD) + 1):
            m = elig & (sec_nat == s)
            rd_nat[m] = RD[s - 1] * (p[m] / p[m].sum() if p[m].sum() > 0 else 1.0 / m.sum())
        rd_lo = rd_hi = rd_sin_tope = rd_nat
        modo_rural, rd_g = "hrsl", None
        log("Sin grilla: la población rural dispersa se reparte por HRSL dentro de cada sección rural.")
    rej_rd, _ = agregar_conservando(rd_nat, tr)
    rej_rd_lo, _ = agregar_conservando(rd_lo, tr)
    rej_rd_hi, _ = agregar_conservando(rd_hi, tr)
    rej_rd_st, _ = agregar_conservando(rd_sin_tope, tr)
    rej_rd_alt = {m_: agregar_conservando(x_, tr)[0] for m_, x_ in rd_alt.items()}
    rej_rd_alt[modo_rural] = rej_rd
    rej_rd_g = agregar_conservando(rd_g, tr)[0] if rd_g is not None else None
    rd_fuera = float(RD.sum() - rej_rd.sum())

    # 4. anclaje 2026 por clase
    f_cab = d26["cabecera"] / P1
    f_res = d26["resto"] / float(T_sec.sum())
    base18 = rej_c1 + rej_c2 + rej_rd
    pob = rej_c1 * f_cab + (rej_c2 + rej_rd) * f_res
    publicar = m_mun | (base18 > 0)
    derrame = float(base18[~m_mun].sum())
    pob[~publicar] = np.nan
    base18_pub = np.where(publicar, base18, np.nan)
    np.save(os.path.join(comun.REJILLA_NPY, "poblacion_cnpv2018.npy"), base18_pub.astype(np.float32))
    assert np.nanmax(pob) < 655.0, "valor fuera del rango codificable con escala 0,01"
    anclaje = {
        "anio": ANIO_ANCLA, "metodo": "censo_2018_por_manzana_y_seccion_rural",
        "dane_cabecera": d26["cabecera"], "dane_resto": d26["resto"], "dane_total": d26["total"],
        "base_cabecera_cnpv2018": P1, "base_resto_cnpv2018": float(T_sec.sum()),
        "factor_cabecera": round(f_cab, 5), "factor_resto": round(f_res, 5),
        "factor_cabecera_omision_y_lea_2018": round(dane["serie"][2018]["cabecera"] / P1, 5),
        "factor_cabecera_crecimiento_2018_2026": round(d26["cabecera"] / dane["serie"][2018]["cabecera"], 5),
        "factor_resto_omision_2018": round(dane["serie"][2018]["resto"] / float(T_sec.sum()), 5),
        "factor_resto_crecimiento_2018_2026": round(d26["resto"] / dane["serie"][2018]["resto"], 5),
        "publicado_por_clase": {"cabecera_manzanas_clase_1": round(float(rej_c1.sum()) * f_cab, 1),
                                "resto_clases_2_y_3_dentro_de_la_rejilla": round(float(rej_c2.sum() + rej_rd.sum()) * f_res, 1)},
        "publicado_en_celdas_de_la_cabecera_mgn": round(float(np.nansum(pob[m_cab])), 1),
        "cabecera_clase_1_en_celdas_con_centro_fuera_de_la_cabecera_mgn": round(float(rej_c1[~m_cab].sum()) * f_cab, 1),
        "publicado_resto_dentro_de_la_rejilla": round(float(np.nansum(pob[m_resto])), 1),
        "publicado_total_rejilla": round(float(np.nansum(pob)), 1),
        "resto_fuera_de_la_rejilla_2026": round((P2 - float(rej_c2.sum()) + rd_fuera) * f_res, 1),
        "poblacion_en_celdas_con_centro_fuera_del_municipio_2018": round(derrame, 1),
    }
    log("Anclaje:", json.dumps(anclaje, ensure_ascii=False))

    # versiones auxiliares y método anterior (para comparar)
    def anclar_clases(q_nat, q_rej):
        s_cab = float(q_rej[m_cab].sum())
        s_resto = float(q_nat[m_mun_nat].sum()) - s_cab
        out = np.full(FORMA, np.nan)
        out[m_cab] = q_rej[m_cab] * d26["cabecera"] / s_cab
        out[m_resto] = q_rej[m_resto] * d26["resto"] / s_resto
        return out, {"factor_cabecera": round(d26["cabecera"] / s_cab, 5), "factor_resto": round(d26["resto"] / s_resto, 5)}
    pob_dos, inf_dos = anclar_clases(p, hrsl)
    np.save(os.path.join(comun.REJILLA_NPY, "poblacion_dos_clases.npy"), pob_dos.astype(np.float32))
    pob_old, inf_old = (anclar_clases(q_old_nat, agregar_conservando(q_old_nat, tr)[0]) if q_old_nat is not None else (None, None))
    anclaje["auxiliares"] = {"dos_clases_hrsl": inf_dos, "grilla_1km_y_hrsl_version_anterior": inf_old}

    # 5. validación
    valid = {}
    # 5a. grilla vs manzanas en cuadros de 1 km enteramente urbanos (dos productos DANE del mismo censo)
    if grilla:
        cel_cab = np.array([c.intersection(cab).area / c.area for c in celdas])
        u = cel_cab >= 0.95
        if u.sum() >= 5:
            valid["grilla_vs_manzanas_1km_urbano"] = {
                "cuadros": int(u.sum()), "personas_grilla": int(G_raw[u].sum()), "personas_manzanas": round(float(M_cel[u].sum())),
                "pearson": round(float(pearsonr(G_raw[u], M_cel[u])[0]), 4),
                "indice_disimilitud": round(indice_disimilitud(G_raw[u], M_cel[u]), 4),
                "nota": "Cuadros de 1 km con ≥ 95 % de su área en la cabecera MGN. Las diferencias vienen de la anonimización "
                        "(manzanas unidas dentro de su sección) y de la asignación de viviendas a cuadros."}
    # 5b. métodos basados en HRSL frente al censo por manzana (cabecera, reparto relativo)
    ref = rej_c1
    metodos = {"hrsl_sin_escalar_o_dos_clases": hrsl}
    if pob_old is not None:
        metodos["grilla_1km_y_hrsl_version_anterior"] = np.nan_to_num(pob_old)
    escalas = {"250_m": 9, "500_m": 18, "1_km": 36}
    valid["metodos_hrsl_vs_censo_por_manzana"] = {
        nom: {esc: comparar_por_bloques(ref, X, m_cab, bb) for esc, bb in escalas.items()} for nom, X in metodos.items()}
    valid["sensibilidad_reparto_dentro_de_manzana"] = {
        "celda_28_m": {"indice_disimilitud": round(indice_disimilitud(rej_c1[m_cab], rej_c1_unif[m_cab]), 4)},
        **{esc: {"indice_disimilitud": comparar_por_bloques(rej_c1_unif, rej_c1, m_cab, bb)["indice_disimilitud"]}
           for esc, bb in escalas.items()},
        "nota": "Publicado (en manzanas ≥ 2 ha: HRSL si cubre ≥ 50 % de lo construido según WorldCover 2021, si no la unión "
                "de ambos) frente a reparto uniforme en todas las manzanas."}
    valid["sensibilidad_regla_manzanas_grandes"] = {
        "celda_28_m": {"indice_disimilitud": round(indice_disimilitud(rej_c1[m_cab] + 1e-12, rej_c1_solo_hrsl[m_cab] + 1e-12), 4)},
        **{esc: {"indice_disimilitud": comparar_por_bloques(rej_c1_solo_hrsl, rej_c1, m_cab, bb)["indice_disimilitud"]}
           for esc, bb in escalas.items()},
        "nota": "Publicado frente a la regla anterior (solo edificaciones HRSL 2016 en toda manzana ≥ 2 ha con alguna)."}
    if grilla:
        rur = ~m_cab & m_mun
        valid["sensibilidad_imputacion_celdas_suprimidas"] = {
            "indice_disimilitud_sin_imputar_vs_publicado": round(indice_disimilitud(rej_rd[rur] + 1e-12, rej_rd_lo[rur] + 1e-12), 4),
            "indice_disimilitud_maximo_vs_publicado": round(indice_disimilitud(rej_rd[rur] + 1e-12, rej_rd_hi[rur] + 1e-12), 4),
            "indice_disimilitud_razones_de_la_grilla_vs_publicado":
                round(indice_disimilitud(rej_rd[rur] + 1e-12, rej_rd_g[rur] + 1e-12), 4) if rej_rd_g is not None else None,
            "nota": "Reparto rural disperso dentro de la rejilla: imputación central frente a suprimidas = 0, = 3 viviendas "
                    "ocupadas y con razones de ocupación y personas por vivienda de la propia grilla."}

    # 5c. fuera de la cabecera: desacuerdo entre las señales de construcción usadas DENTRO de cada cuadro de 1 km
    # (HRSL 2016, WorldCover 2021 y su unión). No hay verdad a escala fina en lo rural: el desacuerdo es una cota de la
    # incertidumbre del reparto, y fija la escala mínima recomendada para sumar celdas fuera de la cabecera.
    escalas_r = {"250_m": 9, "500_m": 18, "1_km": 36}
    def capa_con(rd_rej):
        return rej_c1 * f_cab + (rej_c2 + rd_rej) * f_res
    def di_escalas(A, B, m):
        out = {"celda_28_m": round(indice_disimilitud(A[m] + 1e-12, B[m] + 1e-12), 3)}
        for esc, bb in escalas_r.items():
            a, b_ = bloques(np.where(m, A, 0), bb), bloques(np.where(m, B, 0), bb)
            k = (a > 0) | (b_ > 0)
            out[esc] = round(indice_disimilitud(a[k], b_[k]), 3)
        return out
    try:
        construido = np.nan_to_num(comun.cargar_capa("construido").astype(np.float64))
    except Exception as e:  # noqa: BLE001
        construido = None
        log(f"construido.npy no disponible ({e}): sin diagnóstico con WorldCover en la rejilla común")
    if grilla and len(rej_rd_alt) > 1:
        capas_modo = {m_: capa_con(x_) for m_, x_ in rej_rd_alt.items()}
        pares = [(a_, b_) for a_, b_ in (("hrsl", "worldcover"), ("hrsl", "union"), ("worldcover", "union")) if a_ in capas_modo and b_ in capas_modo]
        from rasterio.features import rasterize as _rz
        vias_osm = [LineString([(lo_, la_) for la_, lo_ in parte]) for v_ in comun.osm().get("vias", []) for parte in v_["partes"]]
        m_vias = _rz([(l_.buffer(15 / 111000), 1) for l_ in vias_osm], out_shape=FORMA, transform=T, fill=0,
                     dtype="uint8").astype(bool) if vias_osm else np.zeros(FORMA, bool)
        idx_rej = indices_grilla(grilla, FORMA, T)
        sec_rej = _rz([(g, k + 1) for k, g in enumerate(g_sec)], out_shape=FORMA, transform=T, fill=0, dtype="int32")
        def di_unidades(A, B, ids_):
            m = m_resto & (ids_ > 0)
            a = np.bincount(ids_[m], weights=A[m]); b_ = np.bincount(ids_[m], weights=B[m])
            k = (a > 0) | (b_ > 0)
            return round(indice_disimilitud(a[k], b_[k]), 3)
        diag = {}
        for m_, X in capas_modo.items():
            x = X[m_resto]
            d = {"suma_hab": round(float(x.sum()), 1), "celdas_pobladas": int((x > 0.005).sum()),
                 "max_hab_celda": round(float(x.max()), 2),
                 "hab_a_15_m_o_menos_de_vias_principales_osm": round(float(X[m_resto & m_vias].sum()), 1)}
            if construido is not None:
                cr = construido[m_resto]
                d.update({"frac_poblacion_en_celdas_con_menos_de_10pct_construido": round(float(x[cr < 10].sum() / x.sum()), 3),
                          "celdas_con_50pct_o_mas_construido": int((cr >= 50).sum()),
                          "de_ellas_con_0_hab": int(((cr >= 50) & (x < 0.005)).sum())})
            diag[m_] = d
        # cuadro de 1 km con más población fuera de la cabecera dentro de la rejilla (el caso más influyente)
        k_max = int(np.argmax(np.bincount(idx_rej[m_resto], weights=capas_modo[modo_rural][m_resto], minlength=len(gf) + 1)[1:]))
        mk = (idx_rej == k_max + 1) & m_resto
        cuadro = {"n6_cod": gf[k_max]["properties"].get("n6_cod"), "lat": round(celdas[k_max].centroid.y, 4),
                  "lon": round(celdas[k_max].centroid.x, 4), "celdas_fuera_de_la_cabecera": int(mk.sum())}
        for m_, X in capas_modo.items():
            x = X[mk]
            cuadro[m_] = {"suma_hab": round(float(x.sum()), 1), "celdas_pobladas": int((x > 0.005).sum()),
                          "max_hab_celda": round(float(x.max()), 2)}
            if construido is not None:
                cr, hb_ = construido[mk], hrsl[mk] > 0
                cuadro[m_]["hab_en_celdas_construidas_50pct_sin_hrsl"] = round(float(x[(cr >= 50) & ~hb_].sum()), 1)
                cuadro[m_]["celdas_construidas_50pct_sin_hrsl"] = int(((cr >= 50) & ~hb_).sum())
        valid["fuera_de_la_cabecera_desacuerdo_entre_senales"] = {
            "publicado": modo_rural,
            "indice_disimilitud_entre_repartos": {f"{a_}_vs_{b_}": di_escalas(capas_modo[a_], capas_modo[b_], m_resto) | {
                "cuadro_grilla_dane_1km": di_unidades(capas_modo[a_], capas_modo[b_], idx_rej),
                "seccion_rural": di_unidades(capas_modo[a_], capas_modo[b_], sec_rej)} for a_, b_ in pares},
            "diagnostico_con_worldcover_2021": diag,
            "cuadro_1km_con_mas_poblacion_fuera_de_la_cabecera": cuadro,
            "nota": "Celdas con dato fuera de la cabecera MGN (capa completa: rural disperso, centros poblados y bordes de "
                    "manzanas de la cabecera). Mismo total por cuadro y por sección en los tres repartos; solo cambia el reparto "
                    "dentro de cada cuadro de 1 km. Índice de disimilitud = fracción de la población que habría que mover."}
    # 5d. prueba análoga en la cabecera (donde sí hay verdad: el censo por manzana con reparto uniforme dentro de la
    # manzana): se reparte el total de cada cuadro de 1 km según cada señal y se compara con el censo
    if grilla and construido is not None:
        idx_rej = indices_grilla(grilla, FORMA, T)
        verdad = rej_c1_unif
        def por_cuadro(m, w):
            n_ = len(gf) + 1
            Tk = np.bincount(idx_rej[m], weights=verdad[m], minlength=n_)
            Wk = np.bincount(idx_rej[m], weights=w[m], minlength=n_)
            Nk = np.bincount(idx_rej[m], minlength=n_)
            out = np.zeros(FORMA)
            ok = m & (Wk[idx_rej] > 0)
            out[ok] = (Tk[idx_rej] * w / np.where(Wk[idx_rej] > 0, Wk[idx_rej], 1))[ok]
            un = m & (Wk[idx_rej] == 0)
            out[un] = (Tk[idx_rej] / np.maximum(Nk[idx_rej], 1))[un]
            return out
        def union_rej(m):
            n_ = len(gf) + 1
            Hk_ = np.bincount(idx_rej[m], weights=hrsl[m], minlength=n_)
            Ck_ = np.bincount(idx_rej[m], weights=construido[m], minlength=n_)
            hn_ = np.where(Hk_[idx_rej] > 0, hrsl / np.where(Hk_[idx_rej] > 0, Hk_[idx_rej], 1), 0)
            cn_ = np.where(Ck_[idx_rej] > 0, construido / np.where(Ck_[idx_rej] > 0, Ck_[idx_rej], 1), 0)
            return np.maximum(hn_, cn_)
        area_ = comun.area_celda_m2()
        m_c = m_cab & (idx_rej > 0)
        dens_k = (np.bincount(idx_rej[m_c], weights=verdad[m_c], minlength=len(gf) + 1)
                  / np.maximum(np.bincount(idx_rej[m_c], weights=area_[m_c], minlength=len(gf) + 1), 1) * 1e6)
        prueba = {}
        for nom_m, m in (("cabecera", m_c), ("cuadros_de_baja_densidad_menos_de_6000_hab_km2", m_c & (dens_k[idx_rej] < 6000))):
            res = {"personas_censo_2018": round(float(verdad[m].sum()))}
            for nom_w, w in (("hrsl_2016", hrsl), ("worldcover_2021", construido), ("union", union_rej(m)), ("uniforme", np.ones(FORMA))):
                X = por_cuadro(m, w)
                res[nom_w] = di_escalas(verdad, X, m)
            prueba[nom_m] = res
        prueba["nota"] = ("Análogo urbano del reparto rural: dentro de cada cuadro de 1 km se reparte el total censal de la parte "
                          "de cabecera según cada señal y se compara con el censo por manzana (índice de disimilitud; menor es "
                          "mejor). En la cabecera WorldCover también capta comercio, industria y vías, y en lo rural puede no ver "
                          "viviendas dispersas bajo árboles: es una prueba indirecta, no una validación del reparto rural.")
        valid["prueba_analoga_senales_dentro_del_cuadro_en_la_cabecera"] = prueba
    log("Validación:", json.dumps(valid, ensure_ascii=False))

    # 6. metadatos y publicación
    area_m2 = comun.area_celda_m2()
    a_med = float(area_m2[m_cab].mean())
    pos = pob[np.isfinite(pob) & (pob > 0)]
    vmax = float(np.ceil(np.percentile(pos, 99)))
    frac_mun_en_rejilla = km2(mun.intersection(rejilla)) / km2(mun)
    cab_osm = comun._poligono(comun.osm()["cabecera"])
    v250 = valid["metodos_hrsl_vs_censo_por_manzana"]
    s250 = valid["sensibilidad_reparto_dentro_de_manzana"]
    vieja = v250.get("grilla_1km_y_hrsl_version_anterior")
    des = valid.get("fuera_de_la_cabecera_desacuerdo_entre_senales")
    ana = valid.get("prueba_analoga_senales_dentro_del_cuadro_en_la_cabecera")
    di_hw = (des or {}).get("indice_disimilitud_entre_repartos", {}).get("hrsl_vs_worldcover")
    di_pub = [v_ for k_, v_ in (des or {}).get("indice_disimilitud_entre_repartos", {}).items() if modo_rural in k_.split("_vs_")]
    pct = lambda x: nd(100 * x, 0)  # noqa: E731
    if di_hw:
        esc_resto_txt = (
            f"Fuera de la cabecera (zona rural y periurbana, ≈ {nd(100 * float(np.nansum(pob[m_resto])) / float(np.nansum(pob)), 1)} % "
            "de la población de la capa) el reparto dentro de cada cuadro de 1 km de la grilla DANE depende de qué señal de "
            "construcción se use, y la escala útil es mayor: sumar por sección rural (su total 2018 es exacto; polígonos en "
            "datos/vectores/poblacion.geojson) o por cuadro de la grilla DANE, donde repartir con edificaciones HRSL (2016) o con "
            f"área construida WorldCover (2021) cambia el resultado en ≈ {pct(di_hw['seccion_rural'])} % y "
            f"{pct(di_hw['cuadro_grilla_dane_1km'])} % (índice de disimilitud). En áreas arbitrarias el desacuerdo crece: "
            f"{pct(di_hw['1_km'])} % en bloques de 1 km, {pct(di_hw['500_m'])} % a 500 m y {pct(di_hw['250_m'])} % a 250 m; la "
            f"versión publicada (unión de ambas señales) difiere de cada una en ≈ {pct(min(v_['1_km'] for v_ in di_pub))}–"
            f"{pct(max(v_['1_km'] for v_ in di_pub))} % a 1 km y ≈ {pct(min(v_['250_m'] for v_ in di_pub))}–"
            f"{pct(max(v_['250_m'] for v_ in di_pub))} % a 250 m. Por eso, fuera de la cabecera, áreas de influencia de 1 km o "
            "menos (un tramo de río, una vereda) dan solo un orden de magnitud.")
    else:
        esc_resto_txt = ("Fuera de la cabecera (zona rural y periurbana): sumar por sección rural (total 2018 exacto) o en áreas "
                         "de 1 km o más; dentro de cada cuadro de 1 km el reparto depende de las edificaciones HRSL de 2016.")
    esc_cab_txt = ("En la cabecera: sumar celdas por manzana, barrio, comuna o áreas de influencia de 250 m o más; a esa escala "
                   "el reparto depende poco del supuesto dentro de manzana (índice de disimilitud a 250 m frente a reparto "
                   f"uniforme: {nd(s250['250_m']['indice_disimilitud'], 3)}).")
    c_ = (des or {}).get("cuadro_1km_con_mas_poblacion_fuera_de_la_cabecera")
    sup_txt = (f"{info_grilla['celdas_suprimidas_que_tocan_cartago']} cuadros de 1 km que tocan Cartago tienen viviendas pero "
               "personas = 0 en la grilla; se tratan como dato suprimido por reserva estadística (supuesto inferido de los datos: "
               "ninguna celda publicada tiene menos de 4 viviendas ocupadas) y se imputan con "
               f"{nd(info_grilla['personas_imputadas_en_celdas_que_tocan_cartago_ponderadas_por_area'], 0)} personas (rango 0–"
               f"{nd(info_grilla['rango_imputacion_ponderado_por_area'][1], 0)}): esperanza de las viviendas ocupadas condicionada "
               f"a ≤ 3 (binomial, tasa de ocupación rural {nd(100 * p_ocu, 1)} %) × {nd(ppo, 2)} personas por vivienda ocupada."
               ) if info_grilla else "Sin grilla: la población rural dispersa se repartió por HRSL dentro de cada sección."
    fin_ = np.isfinite(pob)
    perdida_png = float(pob[fin_].sum() - (np.round(pob[fin_] / 0.01) * 0.01).sum())
    bajo_umbral = int(((pob > 0) & (pob < 0.005)).sum())
    titulo = f"Población estimada por celda ({ANIO_ANCLA})"
    descripcion = (f"Cuántas personas viven aproximadamente en cada cuadro de unos 28 × 28 m del municipio de Cartago en "
                   f"{ANIO_ANCLA}. Parte del censo 2018 del DANE contado por manzana (y por sección en la zona rural) y lo lleva a "
                   f"{ANIO_ANCLA} con las proyecciones oficiales; es una estimación, no un conteo casa a casa.")
    periodo = (f"Totales: proyección DANE {ANIO_ANCLA} (PPED, {dane['actualizado'] or 'actualizada el 30-jul-2025'}). "
               "Reparto espacial: CNPV 2018 por manzana y sección rural; dentro de cada cuadro rural de 1 km y de manzanas "
               "≥ 2 ha, edificaciones HRSL (imágenes 2016) y área construida ESA WorldCover (2021).")
    procesamiento = [
        f"Personas en hogares particulares del CNPV 2018 por manzana (MGN 2018 integrado, {len(fm)} manzanas de Cartago): "
        f"cabecera {n0(P1)} y centros poblados {n0(P2)}; por sección rural: {len(sec_feats)} secciones con {n0(T_sec.sum())} "
        f"personas (centros poblados incluidos). Coinciden con la ficha CNPV 2018: cabecera {n0(ficha['cabecera_hp_2018']) if ficha else '—'}, "
        f"centros poblados y rural disperso {n0(ficha['resto_hp_2018']) if ficha else '—'}.",
        "Cada manzana reparte sus personas de manera uniforme sobre su polígono, calculado en una sub-rejilla de 5 × 5 por celda "
        f"(≈ 5,5 m). En las manzanas de 2 ha o más (supuesto: en manzanas grandes la gente vive donde hay construcciones): "
        f"solo sobre las edificaciones HRSL de 2016 si cubren al menos el {nd(100 * COBERTURA_HRSL_MIN, 0)} % del área construida "
        f"según WorldCover 2021 dentro de la manzana ({info_m1['manzanas_repartidas_sobre_edificaciones_hrsl'] + info_m2['manzanas_repartidas_sobre_edificaciones_hrsl']} "
        f"manzanas); si no la cubren, sobre la unión de ambas ({info_m1['manzanas_repartidas_sobre_union_hrsl_worldcover'] + info_m2['manzanas_repartidas_sobre_union_hrsl_worldcover']} "
        f"manzanas, {n0(info_m1['personas_en_manzanas_con_union'] + info_m2['personas_en_manzanas_con_union'])} personas), y sin "
        "ninguna de las dos, uniforme. Manzanas sin ninguna sub-celda: a la celda de su punto interior.",
        f"Población rural dispersa de cada sección = personas de la sección − personas de sus manzanas de centro poblado "
        f"({n0(RD.sum())} en total). Dentro de la sección (fuera de las zonas urbanas) se reparte en proporción a las personas "
        "rurales de cada cuadro de 1 km de la grilla DANE y, dentro del cuadro, según la unión de dos señales de construcción: "
        "edificaciones HRSL (2016) y área construida WorldCover (2021); el peso de cada píxel es el mayor de sus dos "
        "participaciones normalizadas en el cuadro, así que recibe población si cualquiera de las dos ve construcción (en "
        "cuadros que la ventana de WorldCover no cubre, solo HRSL). Personas rurales "
        "del cuadro = grilla − personas de TODAS las manzanas censales del cuadro, de Cartago y de municipios vecinos "
        f"({info_rural.get('manzanas_vecinas_en_la_ventana', 0)} manzanas vecinas, p. ej. Puerto Caldas de Pereira), descontando "
        f"una tolerancia de {nd(100 * info_rural.get('tolerancia_ruido_tau', 0), 1)} % de las personas en manzanas (la mayor "
        f"diferencia entre grilla y manzanas en los {info_rural.get('cuadros_urbanos_para_tolerancia', 0)} cuadros enteramente "
        "urbanos, que es ruido de agregación y no población rural). " + sup_txt + (
            f" Contraste parcialmente independiente (la imputación usa la ocupación y las personas por vivienda de las secciones "
            f"censales): la grilla así tratada implica {n0(info_grilla['rural_implicito_grilla_menos_manzanas_en_secciones'])} "
            f"personas rurales dispersas en las secciones de Cartago frente a {n0(RD.sum())} del censo por sección (sin imputar, "
            f"{n0(info_rural['rural_implicito_sin_imputar'])}; imputando con razones de la propia grilla —"
            f"{info_rural['imputacion_con_razones_de_la_grilla']['n_celdas']} celdas rurales publicadas, "
            f"{nd(info_rural['imputacion_con_razones_de_la_grilla']['personas_por_vivienda_ocupada'], 2)} personas por vivienda "
            f"ocupada y {nd(100 * info_rural['imputacion_con_razones_de_la_grilla']['tasa_ocupacion'], 1)} % de ocupación—, "
            f"{n0(info_rural['imputacion_con_razones_de_la_grilla']['rural_implicito_en_secciones'])})." if info_grilla else ""),
        "ESA WorldCover 10 m 2021 v200 (clase 50, «construido»: edificios, vías y otras estructuras) se lee de la ventana local "
        "que guarda scripts/worldcover.py (rejilla común + 0,01°), como fracción construida por promedio de sus píxeles de 10 m: "
        "en la rejilla nativa de HRSL para el reparto rural y en la sub-rejilla de 5,5 m (≥ 50 % construido) para las manzanas. "
        "Solo cambia DÓNDE, dentro de cada cuadro de 1 km o de cada manzana grande, queda la población censada en 2018; no "
        "agrega población por construcciones posteriores al censo.",
        f"HRSL (tesela {TESELA_HRSL}) leído por ventana: {arr.shape[1]}×{arr.shape[0]} píxeles de 1\". Su agregación a la rejilla "
        f"conserva la masa (Δ {nd(conserv['dif_vs_masa_fuente_pct'], 4)} % frente a la masa fuente; rasterio Resampling.sum "
        f"Δ {nd(conserv['dif_rasterio_sum_vs_exacto_pct'], 2)} %).",
        f"Anclaje a {ANIO_ANCLA}: cabecera × {nd(f_cab, 4)} = {n0(d26['cabecera'])} / {n0(P1)}; centros poblados y rural "
        f"disperso × {nd(f_res, 4)} = {n0(d26['resto'])} / {n0(T_sec.sum())}. Cada factor combina el ajuste del DANE por "
        f"omisión censal (y, en la cabecera, las {n0(ficha['lea_2018']) if ficha else '1.408'} personas en lugares especiales de "
        f"alojamiento, ≈ {nd(100 * ficha['lea_2018'] / ficha['cabecera_2018'], 1) if ficha else '1,2'} % de la cabecera censada) para 2018 (× {nd(anclaje['factor_cabecera_omision_y_lea_2018'], 3)} y × "
        f"{nd(anclaje['factor_resto_omision_2018'], 3)}) y el crecimiento proyectado 2018→{ANIO_ANCLA} "
        f"(× {nd(anclaje['factor_cabecera_crecimiento_2018_2026'], 3)} y × {nd(anclaje['factor_resto_crecimiento_2018_2026'], 3)}).",
        "Fuera del municipio de Cartago (MGN 2018) la capa queda sin dato. Versiones auxiliares no publicadas en "
        "fuentes/rejilla/: poblacion_cnpv2018.npy (censo 2018 sin proyectar), poblacion_hrsl.npy (HRSL 2020 sin escalar) y "
        "poblacion_dos_clases.npy (HRSL escalado por clase al DANE 2026, método prescrito originalmente).",
        "Codificación PNG de 16 bits con escala 0,01 hab/celda (comun.guardar_capa).",
    ]
    limitaciones = [
        "Escala de uso. " + esc_cab_txt + " " + esc_resto_txt,
        f"La base es el censo de 2018: se supone que todas las manzanas de la cabecera crecieron en la misma proporción hasta "
        f"{ANIO_ANCLA}. Urbanizaciones habitadas después de 2018 no suman población (fuera de la cabecera, WorldCover 2021 puede "
        "ubicar en ellas parte de la población rural censada en 2018, sin aumentarla) y las que se vaciaron conservan su "
        "población de 2018.",
        f"Las {n0(ficha['lea_2018']) if ficha else '1.408'} personas censadas en lugares especiales de alojamiento (LEA: centros "
        "penitenciarios, cuarteles y estaciones de policía, hogares de adultos mayores, conventos, internados, según el DANE) "
        "no tienen ubicación pública: el DANE las publica solo por clase. El factor de la cabecera las reparte en proporción a "
        "la población de los hogares, así que las celdas de esas instituciones quedan subestimadas y el resto levemente sobreestimado.",
        "Reserva estadística del DANE: las manzanas con 1 a 3 viviendas ocupadas se unieron a otra manzana de su misma sección "
        f"urbana ({totales_censo['secciones_urbanas']['n']} secciones en la cabecera; diagonal mediana "
        f"{n0(totales_censo['secciones_urbanas']['diagonal_mediana_m'])} m), así que unas pocas personas aparecen corridas dentro "
        f"de su sección. Las {totales_censo['secciones_urbanas']['manzanas_cabecera_con_viviendas_y_0_personas']} manzanas de la "
        "cabecera con viviendas y 0 personas pueden ser manzanas así o manzanas sin viviendas ocupadas.",
        "Dentro de cada manzana el reparto es uniforme (o, en manzanas de 2 ha o más, sobre lo construido según HRSL 2016 y "
        "WorldCover 2021): la cifra de una celda de 28 m sola es incierta y las calles, parques y zonas sin manzana censal "
        "valen 0 o casi 0.",
        "La zona rural es menos precisa que la urbana: el total de cada sección rural es exacto (2018), pero su reparto interno "
        "usa la grilla de 1 km (con celdas suprimidas imputadas) y, dentro de cada cuadro, dos señales de construcción que no "
        "coinciden: HRSL no ve lo construido después de 2016 y WorldCover marca también vías e instalaciones no residenciales y "
        "puede no ver viviendas dispersas bajo árboles. " + (
            f"En el cuadro con más población fuera de la cabecera ({c_['n6_cod']}, {nd(c_['lat'], 4)}, {nd(c_['lon'], 4)}), "
            f"{c_['hrsl']['celdas_construidas_50pct_sin_hrsl']} celdas con 50 % o más construido según WorldCover no tienen "
            f"edificación HRSL; con solo HRSL valían 0 y con la unión reciben {n0(c_[modo_rural]['hab_en_celdas_construidas_50pct_sin_hrsl'])} hab. "
            if c_ and modo_rural in c_ else "") +
        f"En {info_rural.get('cuadros_que_tocan_cartago_con_reparto_uniforme', 0)} cuadros de 1 km sin ninguna de las dos "
        "señales fuera de lo urbano, la poca población rural asignada (sobre todo la imputada) se reparte de forma uniforme: se "
        "ve como cuadros tenues de 1 km, que significan «unas pocas personas en este km²», no una ocupación homogénea.",
        f"Con la escala de 0,01 hab/celda, las {n0(bajo_umbral)} celdas con menos de 0,005 hab quedan en 0 en la PNG; la "
        f"cuantización resta {nd(perdida_png, 0)} hab al total de la PNG. El .npy conserva los valores sin redondear.",
        "Las clases urbano/rural son las del MGN 2018 del DANE (la cabecera incluye a Zaragoza); si el perímetro urbano vigente "
        "del POT difiere, cambian las celdas que reciben cada factor.",
        f"La rejilla cubre el {nd(100 * frac_mun_en_rejilla, 0)} % del área municipal: la suma de la capa es menor que el total "
        "municipal del DANE (la cabecera sí está completa). Fuera del municipio no se publica valor.",
        "La proyección del DANE parte del CNPV 2018 ajustado por omisión: es una proyección demográfica, no un conteo, y tiene "
        "su propia incertidumbre.",
    ]
    orden = [0, 1, 5, 4]
    limitaciones = [limitaciones[i] for i in orden] + [x for i, x in enumerate(limitaciones) if i not in orden]
    if ana:
        a_ = ana["cuadros_de_baja_densidad_menos_de_6000_hab_km2"]
        limitaciones.append(
            "Por qué la unión de HRSL y WorldCover y no una sola señal (prueba indirecta, en la cabecera, donde el censo por "
            "manzana sirve de verdad): repartiendo el total de cada cuadro de 1 km según cada señal, en los cuadros de baja "
            f"densidad habría que mover el {pct(a_['hrsl_2016']['250_m'])} % de la población en bloques de 250 m con HRSL, el "
            f"{pct(a_['worldcover_2021']['250_m'])} % con WorldCover, el {pct(a_['union']['250_m'])} % con la unión y el "
            f"{pct(a_['uniforme']['250_m'])} % con reparto uniforme. WorldCover gana en la ciudad, pero en lo rural puede no ver "
            "viviendas dispersas que HRSL sí detecta; la unión conserva ambas.")
    if vieja:
        limitaciones.append(
            f"Contraste con métodos basados solo en satélite: frente al censo por manzana, HRSL repartido dentro de cuadros de 1 km "
            f"(versión anterior de esta capa) dejaba el {nd(100 * vieja['250_m']['indice_disimilitud'], 0)} % de la población de la "
            f"cabecera en bloques de 250 m equivocados, y HRSL solo, el "
            f"{nd(100 * v250['hrsl_sin_escalar_o_dos_clases']['250_m']['indice_disimilitud'], 0)} %.")
    fuentes = [
        {"nombre": "DANE — MGN 2018 integrado con el CNPV 2018: manzana censal (personas en hogares particulares)",
         "url": URL_MANZ, "documentacion": URL_INSTRUCTIVO, "geovisor": URL_GEOVISOR_MANZ, "licencia": LIC_DANE,
         "cita": CITA_DANE, "fecha_consulta": fecha_consulta(manz)},
        {"nombre": "DANE — MGN 2018 integrado con el CNPV 2018: sección rural", "url": URL_SECR,
         "documentacion": URL_INSTRUCTIVO, "licencia": LIC_DANE, "cita": CITA_DANE, "fecha_consulta": fecha_consulta(secr)},
        {"nombre": "DANE — Serie municipal de población por área 2018-2042 (PPED)", "url": URL_DANE, "pagina": PAG_DANE,
         "licencia": LIC_DANE, "cita": CITA_DANE + f" (PPED, serie municipal por área 2018-2042; {dane['actualizado']}).",
         "fecha_descarga": fecha_archivo("PPED-AreaMun-2018-2042_VP.xlsx")},
        {"nombre": "DANE — Marco Geoestadístico Nacional 2018 (zona urbana y municipio)", "url": URL_MGN,
         "licencia": LIC_DANE, "cita": CITA_DANE, "fecha_consulta": mgn["fecha"]},
        {"nombre": "DANE — Grilla estadística nivel 6 (1 km), CNPV 2018", "url": URL_GRILLA, "licencia": LIC_DANE,
         "cita": CITA_DANE + " (Geoportal, Grilla DANE)", "fecha_consulta": fecha_consulta(grilla)},
        {"nombre": "DANE — Ficha municipal CNPV 2018, Cartago (76147)", "url": URL_FICHA_CNPV, "licencia": LIC_DANE, "cita": CITA_DANE,
         "fecha_descarga": fecha_archivo("cnpv2018_ficha_76147.pdf")},
        {"nombre": "DANE — Manzanas censales de municipios vecinos dentro de la ventana (se restan de la grilla)", "url": URL_MANZ,
         "licencia": LIC_DANE, "cita": CITA_DANE, "fecha_consulta": fecha_consulta(vecinas)},
        {"nombre": "HRSL v1.5 (Meta Data for Good y CIESIN), población general 2020", "url": URL_VRT,
         "registro": URL_REGISTRO_HRSL, "hdx": URL_HDX_HRSL, "licencia": "CC BY 4.0", "cita": CITA_HRSL.format(hoy=hoy),
         "fecha_lectura_ventana": fecha_archivo("hrsl_ventana.npz")},
        {"nombre": "ESA WorldCover 10 m 2021 v200, clase 50 «construido» (reparto dentro de cuadros rurales de 1 km y de "
                   "manzanas ≥ 2 ha; ventana local de scripts/worldcover.py)", "url": URL_WC,
         "licencia": "CC BY 4.0", "cita": CITA_WC, "atribucion": ATRIB_WC,
         "fecha_lectura_ventana": time.strftime("%Y-%m-%d", time.localtime(os.path.getmtime(WC_VENTANA))) if os.path.exists(WC_VENTANA) else None},
        {"nombre": "OpenStreetMap (límite municipal para la ventana de lectura; comunas, vías principales y puntos de control de las verificaciones)",
         "url": "https://www.openstreetmap.org/copyright", "licencia": "ODbL 1.0", "cita": "© colaboradores de OpenStreetMap"},
    ]
    meta = {
        "titulo": titulo,
        "descripcion": descripcion,
        "unidad": "hab/celda",
        "fuente": {
            "nombre": "DANE (CNPV 2018 por manzana y sección rural, proyecciones PPED por área, MGN 2018, grilla 1 km), "
                      "Meta/CIESIN HRSL v1.5 y ESA WorldCover 2021",
            "url": URL_MANZ,
            "licencia": "DANE: uso y transformación autorizados citando «Fuente: Departamento Administrativo Nacional de "
                        "Estadística: www.dane.gov.co». HRSL y ESA WorldCover: CC BY 4.0.",
            "cita": CITA_DANE + ". " + CITA_HRSL.format(hoy=hoy) + " " + CITA_WC + ". " + ATRIB_WC + ".",
        },
        "fuentes": fuentes,
        "periodo": periodo,
        "resolucion_original_m": round(float(np.sqrt(np.median(area_man[k1])))),
        "resolucion_original_nota": (f"Cabecera: manzana censal (mediana {n0(np.median(area_man[k1]))} m², ≈ "
                                     f"{n0(np.sqrt(np.median(area_man[k1])))} m de lado). Zona rural: sección rural "
                                     f"(≈ {n0(np.mean([km2(g) for g in g_sec]))} km²), grilla de 1 km y HRSL de 30,8 m."),
        "procesamiento": procesamiento,
        "limitaciones": limitaciones,
        "rango_visual": [0, vmax],
        "paleta": ["#fff7ec", "#fdd49e", "#fc8d59", "#d7301f", "#7f0000"],
        "interpretacion": (f"Una celda mide ≈ {n0(a_med)} m²: 1 hab/celda ≈ {n0(1e6 / a_med)} hab/km². No hay umbrales "
                           "normativos; la capa sirve para estimar cuántas personas hay en un área sumando celdas (p. ej., "
                           "población expuesta a calor, inundación o ruido), no para leer celdas sueltas. " + esc_cab_txt + " " +
                           esc_resto_txt + " 0 = celda sin manzana censal habitada (calles, parques, zonas comerciales o "
                           "industriales sin vivienda) o zona rural sin población asignada."),
        "metodo": anclaje["metodo"],
        "anclaje_dane": anclaje,
        "totales_censo_2018": totales_censo,
        "reparto_manzanas": {"cabecera": info_m1, "centros_poblados": info_m2},
        "reparto_rural": info_grilla,
        "validacion": valid,
        "conservacion_agregacion_hrsl": conserv,
    }
    stats = comun.guardar_capa(NOMBRE, pob, escala=0.01, desplazamiento=0.0, meta=meta)

    # 7. verificaciones
    m_osm = comun.mascara_urbana()
    verif = {"totales_censo_2018": totales_censo, "validacion": valid}
    dec = decodificar_png(NOMBRE, 0.01, 0.0)
    npy = comun.cargar_capa(NOMBRE)
    ambos = np.isfinite(dec) & np.isfinite(npy)
    verif["png_vs_npy"] = {
        "mismas_celdas_sin_dato": bool(np.array_equal(np.isfinite(dec), np.isfinite(npy))),
        "max_dif_abs": float(np.max(np.abs(dec[ambos] - npy[ambos]))),
        "suma_png": round(float(np.nansum(dec)), 1), "suma_npy": round(float(np.nansum(npy)), 1)}
    verif["conservacion"] = {
        "cabecera_publicada_vs_dane": [round(float(np.nansum(pob[m_cab])), 1), d26["cabecera"]],
        "manzanas_cabecera_en_rejilla_vs_censo": [round(float(rej_c1.sum()), 1), P1],
        "rural_disperso_en_rejilla_mas_fuera": [round(float(rej_rd.sum()), 1), round(rd_fuera, 1), float(RD.sum())]}

    la, lo = comun.centros_celdas()
    def entorno(capa, lat, lon, radio_m):
        m = np.hypot((la - lat) * 110574, (lo - lon) * 111320 * np.cos(np.radians(lat))) <= radio_m
        x = capa[m]
        ok = np.isfinite(x)
        return {"celdas": int(m.sum()), "con_dato": int(ok.sum()), "suma_hab": round(float(np.nansum(x)), 1),
                "densidad_hab_km2": round(float(np.nansum(x)) / float(area_m2[m & np.isfinite(capa)].sum()) * 1e6) if ok.any() else None,
                "frac_celdas_pobladas": round(float((x[ok] > 0).mean()), 3) if ok.any() else None}
    for nom, (lat, lon) in {"parque_bolivar": (4.7497, -75.9132), "aerodromo": (4.7601, -75.9545)}.items():
        i, j = celda(lat, lon)
        verif[nom] = {"celda": [i, j], "valor_celda": None if not np.isfinite(pob[i, j]) else round(float(pob[i, j]), 2),
                      "radio_300m": entorno(pob, lat, lon, 300),
                      "radio_300m_version_anterior": entorno(pob_old, lat, lon, 300) if pob_old is not None else None}
    pista = comun.osm()["pista"]
    vp = muestrear_linea(pob, LineString([(x, y) for y, x in pista]))
    verif["pista_aerodromo"] = {"celdas_muestreadas": int(vp.size), "media_hab_celda": round(float(np.nanmean(vp)), 3),
                                "frac_en_cero": round(float((vp[np.isfinite(vp)] == 0).mean()), 3)}
    vr = muestrear_linea(pob, comun.lineas_rio())
    ok = np.isfinite(vr)
    verif["rio_la_vieja"] = {
        "celdas_muestreadas": int(vr.size), "con_dato": int(ok.sum()),
        "media_hab_celda": round(float(vr[ok].mean()), 3), "frac_en_cero": round(float((vr[ok] == 0).mean()), 3),
        "media_cabecera_hab_celda": round(float(np.nanmean(pob[m_cab])), 3)}
    # anillo periurbano y zona rural lejana: publicado frente a la versión anterior
    rur_val = {}
    for nom, (d0, d1) in {"anillo_0_500m_fuera_de_la_cabecera": (0, 500), "anillo_500_1000m": (500, 1000),
                          "rural_a_mas_de_1km_de_la_cabecera": (1000, None)}.items():
        mm = m_resto & (anillo(cab, d0, d1, FORMA, T) if d1 else ~mascara(cab.buffer(d0 / 111000), FORMA, T))
        rur_val[nom] = {"celdas": int(mm.sum()), "suma_hab": round(float(np.nansum(pob[mm])), 1),
                        "solo_centros_poblados_y_rural_disperso_hab": round(float((rej_c2 + rej_rd)[mm].sum()) * f_res, 1),
                        "de_manzanas_de_la_cabecera_hab": round(float(rej_c1[mm].sum()) * f_cab, 1),
                        "rural_disperso_sin_tolerancia_hab": round(float(rej_rd_st[mm].sum()) * f_res, 1),
                        "densidad_hab_km2": round(float(np.nansum(pob[mm])) / float(area_m2[mm].sum()) * 1e6, 1),
                        "frac_celdas_pobladas": round(float((pob[mm] > 0).mean()), 4),
                        "version_anterior_suma_hab": round(float(np.nansum(pob_old[mm])), 1) if pob_old is not None else None}
    verif["zona_rural"] = rur_val
    frac = float(hrsl[m_cab].sum() / hrsl.sum())
    frac_v = float(np.flipud(hrsl)[m_cab].sum() / hrsl.sum())
    fracm = float(rej_c1[m_cab].sum() / rej_c1.sum())
    i0, _ = celda(4.7497, -75.9132)
    lat_c = comun.NORTE - (i0 + 0.5) * comun.RES
    fila_nat = int((lat_c - tr.f) / tr.e)
    verif["orientacion"] = {
        "frac_masa_hrsl_en_cabecera": round(frac, 4), "frac_si_se_voltea_n_s": round(frac_v, 4),
        "frac_manzanas_cabecera_dentro_de_la_cabecera_mgn": round(fracm, 4),
        "frac_manzanas_si_se_voltea_n_s": round(float(np.flipud(rej_c1)[m_cab].sum() / rej_c1.sum()), 4),
        "lat_centro_fila_parque_bolivar": round(lat_c, 5), "lat_fila_nativa_equivalente": round(tr.f + (fila_nat + 0.5) * tr.e, 5),
        "norte_arriba": bool(frac > frac_v)}

    zonas = comun.comunas()
    mz = comun.mascara_zonas(zonas)
    por_zona = []
    for z in zonas:
        m = mz == z["indice"]
        por_zona.append({"id": z["id"], "nombre": z["nombre"], "geometria": z["fuente"], "celdas": int(m.sum()),
                         "poblacion_hab": round(float(np.nansum(pob[m]))),
                         "cnpv2018_hogares_particulares_hab": round(float(np.nansum(base18_pub[m]))),
                         "hrsl_2020_sin_escalar_hab": round(float(hrsl[m].sum())),
                         "dos_clases_hab": round(float(np.nansum(pob_dos[m]))),
                         "version_anterior_hab": round(float(np.nansum(pob_old[m]))) if pob_old is not None else None,
                         "area_km2": round(float(area_m2[m].sum()) / 1e6, 3),
                         "densidad_hab_km2": round(float(np.nansum(pob[m])) / float(area_m2[m].sum()) * 1e6) if m.any() else None})
    verif["sumas_de_control"] = {
        "comunas_1_a_7": sum(z["poblacion_hab"] for z in por_zona if z["id"].startswith("c")),
        "cabecera_dane_mgn": round(float(np.nansum(pob[m_cab]))),
        "cabecera_osm": round(float(np.nansum(pob[m_osm]))),
        "cabecera_dane_fuera_de_cabecera_osm": round(float(np.nansum(pob[m_cab & ~m_osm]))),
        "area_cabecera_dane_km2": round(km2(cab), 2), "area_cabecera_osm_km2": round(km2(cab_osm), 2)}
    # comparación por comuna de los métodos basados en HRSL con el censo por manzana (comunas 1-7)
    ids_c = [z for z in por_zona if z["id"].startswith("c")]
    tref = np.array([z["cnpv2018_hogares_particulares_hab"] for z in ids_c], float)
    valid["por_comuna"] = {nom: {"indice_disimilitud": round(indice_disimilitud(tref, np.array([z[k] for z in ids_c], float)), 3)}
                           for nom, k in (("hrsl_sin_escalar", "hrsl_2020_sin_escalar_hab"), ("dos_clases", "dos_clases_hab"),
                                          ("version_anterior", "version_anterior_hab")) if ids_c[0][k] is not None}
    # cuadros de 1 km con mayor diferencia (ámbitos y años comparables: todo en 2018 o todo en 2026)
    if grilla:
        filas = []
        for k, (c, f) in enumerate(zip(celdas, gf)):
            if f["properties"].get("mpio_cod") != MPIO or not rejilla.contains(c) or cel_mun[k] < 0.95:
                continue
            mk = mascara(c, FORMA, T)
            filas.append({"lat": round(c.centroid.y, 4), "lon": round(c.centroid.x, 4),
                          "grilla_cnpv2018": int(G_raw[k]), "censo_manzanas_y_rural_2018": round(float(np.nansum(base18_pub[mk]))),
                          "publicado_2026": round(float(np.nansum(pob[mk]))), "dos_clases_2026": round(float(np.nansum(pob_dos[mk])))})
        filas.sort(key=lambda r: -abs(r["dos_clases_2026"] - r["publicado_2026"]))
        verif["cuadros_1km_mayor_diferencia_dos_clases"] = filas[:3]

    # estadísticas con la cabecera oficial (MGN) además de la de OSM que calcula comun.guardar_capa
    def resumen(mask):
        x = pob[mask & np.isfinite(pob)]
        return {"min": round(float(x.min()), 4), "p10": round(float(np.percentile(x, 10)), 4), "media": round(float(x.mean()), 4),
                "p50": round(float(np.median(x)), 4), "p90": round(float(np.percentile(x, 90)), 4),
                "max": round(float(x.max()), 4), "celdas": int(x.size), "suma_hab": round(float(x.sum()), 1)}
    ruta_meta = os.path.join(comun.CAPAS, f"{NOMBRE}.json")
    mj = json.load(open(ruta_meta, encoding="utf-8"))
    mj["estadisticas"]["cabecera_dane_mgn"] = resumen(m_cab)
    mj["sumas_hab"] = {"rejilla": round(float(np.nansum(pob)), 1), "cabecera_dane_mgn": round(float(np.nansum(pob[m_cab])), 1),
                       "cabecera_urbana_osm": round(float(np.nansum(pob[m_osm])), 1),
                       "resto_del_municipio_en_la_rejilla": round(float(np.nansum(pob[m_resto])), 1),
                       "celdas_fuera_de_la_cabecera_mgn": round(float(np.nansum(pob[m_resto])), 1),
                       "resto_por_clase": anclaje["publicado_por_clase"]["resto_clases_2_y_3_dentro_de_la_rejilla"]}
    mj["nota_estadisticas"] = (f"'cabecera_urbana' usa el polígono de OpenStreetMap de comun.mascara_urbana ({nd(km2(cab_osm), 2)} km², "
                               f"sin Zaragoza). 'cabecera_dane_mgn' usa la cabecera oficial del MGN 2018 ({nd(km2(cab), 2)} km², "
                               "con Zaragoza), que es la del anclaje a las proyecciones DANE. En 'sumas_hab', "
                               "'celdas_fuera_de_la_cabecera_mgn' (= 'resto_del_municipio_en_la_rejilla', nombre que se conserva "
                               "por compatibilidad) suma las celdas con centro fuera de la cabecera MGN: incluye "
                               f"{nd(anclaje['cabecera_clase_1_en_celdas_con_centro_fuera_de_la_cabecera_mgn'], 1)} hab de "
                               "manzanas de la cabecera que caen en celdas del borde y excluye la parte de la población de "
                               "centros poblados y rural dispersa que cae en celdas con centro dentro de la cabecera. "
                               "'resto_por_clase' es la clase DANE «centros poblados y rural disperso» dentro de la rejilla "
                               "(la cifra comparable con el DANE).")
    mj["validacion"] = valid
    comun.guardar_json(ruta_meta, mj)


    # 8. serie y vector
    totales = [
        {"fuente": "DANE CNPV 2018 por manzana (MGN integrado)", "anio": 2018, "ambito": "cabecera (personas en hogares particulares)",
         "valor_hab": int(P1), "url": URL_MANZ},
        {"fuente": "DANE CNPV 2018 por manzana (MGN integrado)", "anio": 2018, "ambito": "centros poblados (manzanas)",
         "valor_hab": int(P2), "url": URL_MANZ},
        {"fuente": "DANE CNPV 2018 por sección rural (MGN integrado)", "anio": 2018,
         "ambito": "centros poblados y rural disperso (secciones rurales)", "valor_hab": int(T_sec.sum()), "url": URL_SECR},
        {"fuente": "HRSL v1.5 (Meta/CIESIN)", "anio": ANIO_HRSL, "ambito": "municipio de Cartago (MGN 2018, completo)",
         "valor_hab": round(h_mun), "url": URL_VRT, "nota": "suma a resolución nativa 1\""},
        {"fuente": "HRSL v1.5 (Meta/CIESIN)", "anio": ANIO_HRSL, "ambito": "cabecera (MGN 2018)",
         "valor_hab": round(float(hrsl[m_cab].sum())), "url": URL_VRT, "nota": f"celdas de la rejilla; a resolución nativa {n0(h_cab_nat)}"},
        {"fuente": "HRSL v1.5 (Meta/CIESIN)", "anio": ANIO_HRSL, "ambito": "resto del municipio (municipio − cabecera)",
         "valor_hab": round(h_mun - float(hrsl[m_cab].sum())), "url": URL_VRT},
        {"fuente": "HRSL v1.5 (Meta/CIESIN)", "anio": ANIO_HRSL, "ambito": "rejilla de análisis completa (incluye otros municipios)",
         "valor_hab": round(suma_rejilla), "url": URL_VRT},
        {"fuente": "HRSL v1.5 (Meta/CIESIN)", "anio": ANIO_HRSL, "ambito": "parte del municipio fuera de la rejilla",
         "valor_hab": round(h_mun_fuera), "url": URL_VRT},
    ]
    for amb, k in (("cabecera (MGN 2018)", "cabecera"), ("municipio de Cartago (MGN 2018)", "municipio"),
                   ("rejilla de análisis completa", "rejilla")):
        if wp.get(k) is not None:
            totales.append({"fuente": "WorldPop wpgppop (100 m, no restringido)", "anio": ANIO_WP, "ambito": amb,
                            "valor_hab": round(wp[k]), "url": f"{URL_WP}?dataset=wpgppop&year={ANIO_WP}",
                            "nota": "API de estadísticas por polígono (POST con el GeoJSON del ámbito)"})
    for anio in (2018, 2020, 2025, 2026):
        if anio in dane["serie"]:
            for amb, k in (("cabecera municipal", "cabecera"), ("centros poblados y rural disperso", "resto"),
                           ("total municipal", "total")):
                totales.append({"fuente": "DANE PPED (proyección municipal por área)", "anio": anio, "ambito": amb,
                                "valor_hab": dane["serie"][anio][k], "url": URL_DANE})
    if ficha:
        for amb, k in (("cabecera, total censado", "cabecera_2018"), ("cabecera, hogares particulares", "cabecera_hp_2018"),
                       ("centros poblados y rural disperso", "resto_2018"), ("total censado", "total_censado_2018"),
                       ("lugares especiales de alojamiento (LEA)", "lea_2018")):
            totales.append({"fuente": "DANE CNPV 2018, ficha municipal (personas efectivamente censadas, sin ajuste de omisión)",
                            "anio": 2018, "ambito": amb, "valor_hab": ficha[k], "url": URL_FICHA_CNPV})
    if grilla_tot:
        totales.append({"fuente": "DANE Grilla estadística 1 km (CNPV 2018)", "anio": 2018,
                        "ambito": f"{grilla_tot['celdas']} celdas de 1 km codificadas como Cartago (celdas suprimidas en 0)",
                        "valor_hab": grilla_tot["personas"], "url": URL_GRILLA})
    if info_grilla:
        totales.append({"fuente": "DANE Grilla estadística 1 km (CNPV 2018)", "anio": 2018,
                        "ambito": "repartida con HRSL dentro del polígono municipal (MGN 2018), celdas suprimidas en 0",
                        "valor_hab": round(info_grilla["grilla_repartida_con_hrsl_dentro_del_municipio"]), "url": URL_GRILLA})
    totales.append({"fuente": "Capa 'poblacion' (este producto)", "anio": ANIO_ANCLA,
                    "ambito": "celdas con dato dentro de la rejilla", "valor_hab": round(float(np.nansum(pob))), "url": None})
    fuentes_serie = fuentes + [
        {"nombre": "WorldPop — Global per country 2000-2020 (wpgppop), Colombia 2020", "url": URL_WP_META,
         "licencia": "CC BY 4.0", "cita": CITA_WP}] if wp else list(fuentes)
    ruta_z = os.path.join(comun.DATOS, "zaragoza.json")
    if os.path.exists(ruta_z):
        zp = json.load(open(ruta_z, encoding="utf-8"))["properties"]
        fuentes_serie.append({"nombre": "Zona 'Zaragoza' de por_zona: " + zp.get("fuente", ""), "url": "https://doi.org/10.5281/zenodo.7254221",
                              "licencia": zp.get("licencia", "ESA WorldCover: CC BY 4.0"), "cita": zp.get("cita"),
                              "atribucion": zp.get("atribucion")})
    fuentes_serie.append({"nombre": "Definición de lugar especial de alojamiento (LEA), DANE", "url": URL_LEA,
                          "licencia": LIC_DANE, "cita": CITA_DANE})
    serie = {
        "titulo": "Población de Cartago: totales por fuente",
        "unidad": "habitantes",
        "descripcion": "Cifras de población de Cartago según el DANE (proyecciones, censo 2018 por manzana y por sección), HRSL "
                       "(Meta/CIESIN) y WorldPop, y totales de la capa 'poblacion' por comuna (comunas de OpenStreetMap).",
        "atribucion": "Fuente: Departamento Administrativo Nacional de Estadística: www.dane.gov.co · HRSL © Meta y CIESIN (CC BY 4.0) "
                      "· © colaboradores de OpenStreetMap (ODbL) · © ESA WorldCover project 2021 (CC BY 4.0)",
        "fuentes": fuentes_serie,
        "totales": totales,
        "metodo_publicado": anclaje["metodo"],
        "dane_serie_cartago": {"municipio": "Cartago (76147), Valle del Cauca", "actualizado": dane["actualizado"],
                               "url": URL_DANE, "pagina": PAG_DANE, "anios": {str(k): v for k, v in dane["serie"].items()}},
        "anclaje": anclaje,
        "validacion": valid,
        "por_zona": por_zona,
        "nota_por_zona": ("'poblacion_hab' es la capa publicada (2026). 'cnpv2018_hogares_particulares_hab' es el censo 2018 por "
                          "manzana sin proyectar. Zaragoza es parte de la cabecera municipal para el DANE (MGN 2018), aunque "
                          "comun.comunas() la rotule como corregimiento."),
        "limitaciones": meta["limitaciones"],
        "fecha_proceso": hoy,
    }
    comun.guardar_json(os.path.join(comun.DATOS, "series", f"{NOMBRE}.json"), serie)

    def redondear(g, tol=0.00003):   # ≈ 3 m: solo para el vector publicado, no para los cálculos
        return json.loads(json.dumps(mapping(g.simplify(tol, preserve_topology=True))), parse_float=lambda s: round(float(s), 6))
    feats_v = [
        {"type": "Feature", "geometry": redondear(cab), "properties": {
            "clase": "cabecera", "nombre": "Cabecera municipal de Cartago (MGN 2018, incluye Zaragoza)",
            "area_km2": round(km2(cab), 3), "cnpv2018_hogares_particulares": int(P1), "poblacion_dane": d26["cabecera"],
            "anio": ANIO_ANCLA, "factor": anclaje["factor_cabecera"], "fuente": "DANE, MGN 2018 y CNPV 2018 por manzana", "url": URL_MGN}},
        {"type": "Feature", "geometry": redondear(mun), "properties": {
            "clase": "municipio", "nombre": "Municipio de Cartago (MGN 2018)", "area_km2": round(km2(mun), 2),
            "poblacion_dane": d26["total"], "poblacion_resto_dane": d26["resto"], "anio": ANIO_ANCLA,
            "fuente": "DANE, MGN 2018", "url": URL_MGN}}]
    for k, f in enumerate(sec_feats):
        feats_v.append({"type": "Feature", "geometry": redondear(g_sec[k]), "properties": {
            "clase": "seccion_rural", "codigo": cod_sec[k], "area_km2": round(km2(g_sec[k]), 2),
            "cnpv2018_personas": int(T_sec[k]), "cnpv2018_centros_poblados": int(C_sec[k]), "cnpv2018_rural_disperso": int(RD[k]),
            "estimacion_2026": round(float(T_sec[k]) * f_res), "anio": ANIO_ANCLA, "factor": anclaje["factor_resto"],
            "fuente": "DANE, MGN 2018 integrado con el CNPV 2018 (sección rural)", "url": URL_SECR}})
    vec = {"type": "FeatureCollection", "features": feats_v,
           "atribucion": "Fuente: Departamento Administrativo Nacional de Estadística: www.dane.gov.co",
           "nota": "Geometrías del MGN 2018 simplificadas a ≈ 3 m para la web; los cálculos usan las originales."}
    comun.guardar_json(os.path.join(comun.DATOS, "vectores", f"{NOMBRE}.geojson"), vec)

    # 9. resumen
    log("\n================ RESUMEN poblacion ================")
    log(f"HRSL: {info_hrsl}; ventana {arr.shape}; píxeles poblados {int(np.isfinite(arr).sum())}")
    log(f"Cabecera MGN {nd(km2(cab), 2)} km² (OSM {nd(km2(cab_osm), 2)}); municipio {nd(km2(mun), 1)} km²; "
        f"rejilla cubre {nd(100 * frac_mun_en_rejilla, 1)} % del municipio")
    log(f"CNPV 2018: cabecera {n0(P1)} | centros poblados {n0(P2)} | rural disperso {n0(RD.sum())} | secciones {n0(T_sec.sum())}")
    log(f"HRSL 2020: municipio {n0(h_mun)} | rejilla {n0(suma_rejilla)} | municipio fuera de la rejilla {n0(h_mun_fuera)}")
    log(f"WorldPop {ANIO_WP}: {wp}")
    for anio in (2018, 2020, 2025, 2026):
        log(f"DANE {anio}: {dane['serie'].get(anio)}")
    log(f"DANE: {dane['actualizado']}; ficha CNPV 2018: {ficha}; grilla Cartago: {grilla_tot}")
    log(f"Estadísticas capa: {json.dumps(stats, ensure_ascii=False)}")
    log(f"Rango visual: [0, {vmax}]; máximo {nd(float(np.nanmax(pob)), 2)}")
    log("Verificaciones:")
    for k, v in verif.items():
        log(f"  {k}: {json.dumps(v, ensure_ascii=False)}")
    log("Por zona:")
    for z in por_zona:
        log(f"  {z['nombre']}: {n0(z['poblacion_hab'])} hab 2026 (censo 2018 {n0(z['cnpv2018_hogares_particulares_hab'])}), "
            f"{z['area_km2']} km², {z['densidad_hab_km2']} hab/km² (HRSL {n0(z['hrsl_2020_sin_escalar_hab'])}; dos clases "
            f"{z['dos_clases_hab']}; versión anterior {z['version_anterior_hab']})")
    comun.guardar_json(os.path.join(DIR, "verificaciones.json"), verif)


if __name__ == "__main__":
    main()
